from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent

INK = "1F2937"
BLUE = "1F4E79"
GRAY = "F3F4F6"
MID_GRAY = "D1D5DB"
LIGHT_BLUE = "EAF3FA"
PALE_YELLOW = "FFF7D6"


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def borders(cell, color: str = "BFC7D1") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_text(cell, text: str, bold: bool = False, color: str = INK, size: float = 10.0) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders(cell)


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.08

    for name, size in [("Title", 22), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)


def add_header_footer(doc: Document, report_name: str) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = report_name
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("6B7280")

    footer = section.footer.paragraphs[0]
    footer.text = "HairSalon Booking | raport projektowy"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor.from_string("6B7280")


def add_identification(doc: Document, report_kind: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    r = p.add_run(report_kind)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("HairSalon Booking - backend aplikacji salonu fryzjerskiego")
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string("4B5563")

    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(20)
    heading.paragraph_format.space_after = Pt(6)
    run = heading.add_run("Identyfikacja raportu")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    table = doc.add_table(rows=7, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Nazwa przedmiotu", "[wpisz nazwe przedmiotu]"),
        ("Grupa", "[wpisz grupe]"),
        ("Rok akademicki", "[wpisz rok akademicki]"),
        ("Student 1", "Wiktor Gochnio, nr indeksu: [wpisz numer indeksu]"),
        ("Student 2", "[imie i nazwisko], nr indeksu: [wpisz numer indeksu]"),
        ("Prowadzacy", "[wpisz prowadzacego, jesli wymagane]"),
        ("Data przygotowania", "18.05.2026"),
    ]
    for row, (k, v) in zip(table.rows, rows):
        set_text(row.cells[0], k, True, BLUE)
        shade(row.cells[0], LIGHT_BLUE)
        set_text(row.cells[1], v)
    doc.add_page_break()


def add_note(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, PALE_YELLOW)
    set_text(cell, "Uwaga do uzupelnienia: " + text, False, INK, 9.5)
    doc.add_paragraph()


def add_placeholder(doc: Document, caption: str, hint: str, height_lines: int = 7) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, GRAY)
    borders(cell, "9CA3AF")
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("[MIEJSCE NA ZRZUT EKRANU]\n")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    r = p.add_run(hint)
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string("4B5563")
    for _ in range(height_lines):
        p.add_run("\n")
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)
    doc.add_paragraph()


def add_code(doc: Document, code: str, caption: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, "111827")
    borders(cell, "111827")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for i, line in enumerate(code.splitlines()):
        if i:
            p.add_run("\n")
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string("D1FAE5")
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)
    doc.add_paragraph()


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_text(table.rows[0].cells[i], header, True, BLUE)
        shade(table.rows[0].cells[i], LIGHT_BLUE)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_text(cells[i], value)
    doc.add_paragraph()


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_cloud_report() -> Path:
    doc = Document()
    style_doc(doc)
    add_header_footer(doc, "Raport chmurowy - Azure")
    add_identification(doc, "Raport 1: Chmura Azure")

    doc.add_heading("1. Tytul projektu", level=1)
    doc.add_paragraph("HairSalon Booking - backend aplikacji salonu fryzjerskiego do umawiania wizyt.")

    doc.add_heading("2. Cel projektu", level=1)
    doc.add_paragraph(
        "Celem projektu bylo zaprojektowanie i wdrozenie backendu aplikacji salonu fryzjerskiego w chmurze Microsoft Azure. "
        "System ma obslugiwac rezerwacje wizyt, zarzadzanie klientami, fryzjerami i uslugami, a takze pokazywac praktyczne uzycie uslug chmurowych."
    )
    bullets(doc, [
        "Udostepnienie REST API z pelnym CRUD testowanym w Swaggerze.",
        "Przechowywanie danych w Azure Cosmos DB.",
        "Przechowywanie plikow i zdjec w Azure Blob Storage.",
        "Przetwarzanie zdarzen wizyt przez Azure Queue Storage i Azure Functions.",
        "Wysylka potwierdzenia wizyty przez Azure Communication Services Email.",
        "Powiadomienia realtime przez Azure SignalR Service.",
        "Logowanie przez chmure z wykorzystaniem Easy Auth, Google i GitHub.",
    ])

    doc.add_heading("3. Zalozenia projektu", level=1)
    doc.add_heading("3.1 Architektura", level=2)
    doc.add_paragraph(
        "Aplikacja zostala zbudowana jako backend ASP.NET Core uruchomiony w Azure App Service. Logika domenowa jest oddzielona od infrastruktury, a zadania w tle sa realizowane przez kolejke i Azure Function."
    )
    add_code(doc, """Uzytkownik / Swagger / frontend
        |
        v
Azure App Service: booking-api
        |
        +--> Azure Cosmos DB
        +--> Azure Blob Storage
        +--> Azure Queue Storage
        +--> Azure SignalR Service
        |
        v
Azure Function: booking-function
        |
        v
Azure Communication Services Email""", "Schemat 1. Uproszczona architektura chmurowa aplikacji.")

    doc.add_heading("3.2 Wybrane technologie wraz z uzasadnieniem", level=2)
    add_table(doc, ["Technologia / komponent", "Zastosowanie", "Uzasadnienie"], [
        ["ASP.NET Core / C#", "Backend REST API", "Zgodne z wymaganiami projektu, dobra integracja z Visual Studio i Azure."],
        ["Azure App Service", "Hosting API booking-api", "Pozwala uruchomic aplikacje webowa bez recznego zarzadzania serwerem."],
        ["Azure Cosmos DB", "Baza danych NoSQL", "Dobrze przechowuje dokumenty JSON i pasuje do kontenerow customers, barbers, appointments, services, users."],
        ["Azure Blob Storage", "Zdjecia salonu/fryzjerow i pliki podsumowan", "Przeznaczony do plikow binarnych i statycznych."],
        ["Azure Queue Storage", "Kolejka booking-notifications", "Oddziela utworzenie wizyty od wysylki maila."],
        ["Azure Functions", "Serverless wysylajacy maila", "Funkcja uruchamia sie automatycznie po komunikacie w kolejce."],
        ["Azure Communication Services Email", "Mail z potwierdzeniem rezerwacji", "Chmurowa usluga do wysylki emaili."],
        ["Azure SignalR Service", "Powiadomienia realtime", "Pozwala natychmiast pokazac nowa wizyte w podlaczonym kliencie."],
        ["Azure Key Vault", "Bezpieczne sekrety", "Klucze i connection stringi nie powinny byc trzymane w kodzie."],
        ["Azure Monitor / Application Insights", "Logi i diagnostyka", "Pomaga analizowac requesty, kolejki, funkcje i bledy."],
        ["Easy Auth", "Logowanie Google/GitHub", "Szybkie zabezpieczenie App Service bez pisania wlasnego serwera logowania."],
    ])

    doc.add_heading("4. Funkcjonalnosc", level=1)
    doc.add_paragraph("Ponizej opisano glowne funkcjonalnosci aplikacji. W raporcie nalezy uzupelnic je zrzutami ekranow z dzialajacego systemu.")
    bullets(doc, [
        "Klienci: dodawanie, edycja, odczyt, usuwanie, profil klienta i historia wizyt.",
        "Fryzjerzy: dodawanie, edycja, odczyt, usuwanie, specjalizacja, aktywnosc oraz zdjecia.",
        "Uslugi salonu: CRUD uslug, cena, czas trwania i dostepnosc.",
        "Wizyty: tworzenie rezerwacji, walidacja danych, sprawdzanie kolizji terminow, status wizyty.",
        "Konta i role: klient, fryzjer, admin; klient rejestruje sie bez wyboru roli, admin moze nadac role fryzjera.",
        "Email: po utworzeniu wizyty klient otrzymuje potwierdzenie na maila.",
        "SignalR: po utworzeniu wizyty wysylane jest powiadomienie realtime.",
        "Zdjecia: mozliwosc dodawania zdjec salonu i fryzjerow przez Blob Storage.",
    ])
    add_placeholder(doc, "Zrzut 1. Swagger UI z lista kontrolerow.", "Wklej ekran /swagger/index.html z kontrolerami Customers, Hairdressers, Appointments, SalonServices, Auth/Admin.")
    add_placeholder(doc, "Zrzut 2. Tworzenie wizyty w Swaggerze.", "Wklej ekran POST /api/Appointments z odpowiedzia 201 Created.")
    add_placeholder(doc, "Zrzut 3. Mail z potwierdzeniem wizyty.", "Wklej ekran otrzymanej wiadomosci email od DoNotReply@...azurecomm.net.")
    add_placeholder(doc, "Zrzut 4. SignalR test.", "Wklej ekran strony /api/Notifications/signalr-test z komunikatem appointmentBooked.")

    doc.add_heading("5. Przygotowanie projektu", level=1)
    doc.add_heading("5.1 Przygotowanie zasobow Azure", level=2)
    doc.add_paragraph("Projekt wymagal przygotowania grupy zasobow w Azure oraz konfiguracji poszczegolnych uslug.")
    add_table(doc, ["Krok", "Opis"], [
        ["1", "Utworzenie grupy zasobow Booking."],
        ["2", "Utworzenie App Service dla API: booking-api."],
        ["3", "Utworzenie Cosmos DB account booking-db i bazy BookingApkaDB."],
        ["4", "Utworzenie kontenerow Cosmos DB: customers, barbers, services, appointments, users."],
        ["5", "Utworzenie Storage Account bookingfryzjer, kolejki booking-notifications i kontenerow Blob Storage."],
        ["6", "Utworzenie Function App booking-function i powiazanie jej z AzureWebJobsStorage."],
        ["7", "Dodanie Azure SignalR Service i connection stringa do konfiguracji App Service."],
        ["8", "Dodanie Azure Communication Services Email oraz MailFrom jako Email__SenderAddress."],
        ["9", "Dodanie Easy Auth oraz providerow Google/GitHub."],
        ["10", "Dodanie Key Vault / Monitor / Application Insights do konfiguracji i diagnostyki."],
    ])
    add_placeholder(doc, "Zrzut 5. Grupa zasobow Booking w Azure Portal.", "Wklej ekran z Azure Portal pokazujacy liste zasobow projektu.")
    add_placeholder(doc, "Zrzut 6. Cosmos DB Data Explorer.", "Wklej ekran z kontenerami bazy danych i przykladowym dokumentem.")
    add_placeholder(doc, "Zrzut 7. Queue Storage booking-notifications.", "Wklej ekran kolejki lub logow potwierdzajacych przetwarzanie wiadomosci.")

    doc.add_heading("5.2 Przygotowanie kodu", level=2)
    doc.add_paragraph("Kod zostal podzielony na trzy projekty, aby oddzielic API, logike domenowa i funkcje serverless.")
    add_table(doc, ["Projekt", "Rola"], [
        ["HairSalon.Booking.Api", "Kontrolery REST, Swagger, SignalR Hub, integracje z Cosmos DB, Blob Storage i Queue Storage."],
        ["HairSalon.Booking.Core", "Modele domenowe, repozytoria, logika rezerwacji, zdarzenia i wzorce projektowe."],
        ["HairSalon.Booking.Functions", "Azure Function odbierajaca komunikaty z kolejki i wysylajaca email."],
    ])
    add_code(doc, """app.MapControllers();
app.MapHub<BookingNotificationsHub>("/hubs/booking-notifications");""", "Kod 1. Mapowanie kontrolerow i huba SignalR w Program.cs.")
    add_code(doc, """[QueueTrigger("booking-notifications", Connection = "AzureWebJobsStorage")]
public async Task Run(string message, CancellationToken cancellationToken)
{
    var appointment = JsonSerializer.Deserialize<AppointmentBookedMessage>(message, JsonOptions);
    await emailSender.SendAppointmentConfirmationAsync(appointment, cancellationToken);
}""", "Kod 2. Azure Function reagujaca na wiadomosc z kolejki.")
    add_placeholder(doc, "Zrzut 8. Visual Studio - struktura projektu.", "Wklej Solution Explorer z trzema projektami: Api, Core, Functions.")

    doc.add_heading("5.3 Przeplyw kodu widoczny w Visual Studio", level=2)
    doc.add_paragraph(
        "Ta czesc pokazuje, jak kolejne pliki w Visual Studio wspolpracuja z uslugami Azure. "
        "Najlatwiej prezentowac projekt od pliku Program.cs, potem przejsc do kontrolera wizyt, logiki rezerwacji, kolejki oraz funkcji wysylajacej email."
    )
    add_table(doc, ["Etap", "Plik w Visual Studio", "Co pokazac podczas prezentacji"], [
        ["1", "HairSalon.Booking.Api/Program.cs", "Start aplikacji, Swagger, CORS, Key Vault, Application Insights i mapowanie kontrolerow."],
        ["2", "Infrastructure/ServiceCollectionExtensions.cs", "Rejestracja Cosmos DB, Blob Storage, SignalR, repozytoriow i handlerow."],
        ["3", "Controllers/AppointmentsController.cs", "Endpoint POST /api/Appointments wywolywany ze Swaggera."],
        ["4", "Core/Services/AppointmentBookingFacade.cs", "Logika rezerwacji: pobranie danych, walidacja, zapis i publikacja zdarzenia."],
        ["5", "Infrastructure/AzureQueueAppointmentBookedHandler.cs", "Wyslanie komunikatu do Azure Queue Storage."],
        ["6", "Functions/AppointmentBookedQueueFunction.cs", "Serverless Function odbiera komunikat z kolejki."],
        ["7", "Functions/Email/AzureCommunicationEmailSender.cs", "Wyslanie potwierdzenia przez Azure Communication Services Email."],
        ["8", "Infrastructure/SignalRAppointmentBookedHandler.cs", "Powiadomienie realtime przez Azure SignalR Service."],
    ])

    doc.add_heading("5.4 Program.cs - uruchomienie API i polaczenie z Azure", level=2)
    doc.add_paragraph(
        "Program.cs jest pierwszym plikiem, ktory warto pokazac w Visual Studio. "
        "Tutaj aplikacja pobiera konfiguracje z Key Vault, wlacza monitoring, Swaggera, CORS i mapuje endpointy."
    )
    add_code(doc, """var keyVaultUri = builder.Configuration["Azure:KeyVault:VaultUri"];
if (!string.IsNullOrWhiteSpace(keyVaultUri))
{
    builder.Configuration.AddAzureKeyVault(
        new Uri(keyVaultUri),
        new DefaultAzureCredential(),
        new AzureKeyVaultConfigurationOptions());
}

builder.Services.AddApplicationInsightsTelemetry();
builder.Services.AddControllers();
builder.Services.AddEndpointsApiExplorer();
builder.Services.AddSwaggerGen();

app.UseSwagger();
app.UseSwaggerUI();
app.MapControllers();
app.MapHub<BookingNotificationsHub>("/hubs/booking-notifications");""", "Kod 3. Program.cs - konfiguracja Key Vault, monitoringu, Swaggera, kontrolerow i SignalR.")
    doc.add_paragraph(
        "Dzieki temu sekrety nie musza byc zapisane w kodzie. App Service moze odczytywac je z konfiguracji lub Key Vault, a Application Insights zbiera logi i informacje o requestach."
    )
    add_placeholder(doc, "Zrzut 9. Visual Studio - Program.cs.", "Wklej ekran z Program.cs pokazujacy AddAzureKeyVault, AddApplicationInsightsTelemetry, UseSwagger i MapHub.")

    doc.add_heading("5.5 Dependency Injection - gdzie wybierane sa implementacje Azure", level=2)
    doc.add_paragraph(
        "W ServiceCollectionExtensions.cs rejestrowane sa klasy, ktore pozniej sa wstrzykiwane do kontrolerow i serwisow. "
        "Dzieki temu kontrolery nie tworza recznie klientow Cosmos DB, Blob Storage ani SignalR."
    )
    add_code(doc, """services.AddScoped<IPhotoStorageService, BlobPhotoStorageService>();
services.AddScoped<IAppointmentBookingFacade, AppointmentBookingFacade>();
services.AddScoped<IBookingEventPublisher, BookingEventPublisher>();
services.AddScoped<IAppointmentBookedHandler, AzureBlobAppointmentSummaryHandler>();
services.AddScoped<IAppointmentBookedHandler, AzureQueueAppointmentBookedHandler>();
services.AddScoped<IAppointmentBookedHandler, SignalRAppointmentBookedHandler>();

var signalRBuilder = services.AddSignalR();
if (!string.IsNullOrWhiteSpace(options.SignalR.ConnectionString))
{
    signalRBuilder.AddAzureSignalR(options.SignalR.ConnectionString);
}

services.AddSingleton(_ => new CosmosClient(options.Cosmos.ConnectionString));
services.AddScoped(typeof(IBookingRepository<>), typeof(CosmosBookingRepository<>));""", "Kod 4. Rejestracja integracji Azure w kontenerze DI.")
    doc.add_paragraph(
        "Najwazniejsze jest to, ze aplikacja pracuje na interfejsach, np. IBookingRepository albo IAppointmentBookedHandler. "
        "W praktyce oznacza to, ze kontroler wywoluje logike biznesowa, a szczegoly Azure sa schowane w warstwie Infrastructure."
    )

    doc.add_heading("5.6 Endpoint tworzenia wizyty - od Swaggera do logiki domenowej", level=2)
    doc.add_paragraph(
        "Gdy uzytkownik testuje POST /api/Appointments w Swaggerze, request trafia do AppointmentsController. "
        "Kontroler buduje obiekt Appointment i przekazuje go do fasady rezerwacji."
    )
    add_code(doc, """[HttpPost]
public async Task<IActionResult> Create(CreateAppointmentRequest request, CancellationToken cancellationToken)
{
    var appointment = new Appointment
    {
        CustomerId = request.CustomerId,
        HairdresserId = request.HairdresserId,
        SalonServiceId = request.SalonServiceId,
        StartAt = request.StartAt,
        Status = AppointmentStatus.Booked
    };

    var created = await bookingFacade.BookAsync(appointment, cancellationToken);
    return CreatedAtAction(nameof(GetById), new { id = created.id }, created);
}""", "Kod 5. AppointmentsController - endpoint, ktory jest wywolywany ze Swaggera.")
    add_code(doc, """public async Task<Appointment> BookAsync(Appointment appointment, CancellationToken cancellationToken)
{
    var customer = await customers.GetAsync(appointment.CustomerId, cancellationToken);
    var hairdresser = await hairdressers.GetAsync(appointment.HairdresserId, cancellationToken);
    var service = await services.GetAsync(appointment.SalonServiceId, cancellationToken);

    appointment.EndAt = appointment.StartAt.AddMinutes(service.DurationMinutes);
    await EnsureSlotIsFreeAsync(appointment, cancellationToken);

    var created = await appointments.CreateAsync(appointment, cancellationToken);
    await eventPublisher.PublishAppointmentBookedAsync(created, cancellationToken);
    return created;
}""", "Kod 6. AppointmentBookingFacade - glowna logika rezerwacji wizyty.")
    doc.add_paragraph(
        "Ten fragment jest dobry do omowienia dzialania aplikacji: najpierw sprawdzane sa powiazane dane, potem wyliczana jest godzina zakonczenia, nastepnie system sprawdza kolizje terminu, zapisuje wizyte i uruchamia reakcje po rezerwacji."
    )
    add_placeholder(doc, "Zrzut 10. Visual Studio - AppointmentsController i AppointmentBookingFacade.", "Wklej ekran pokazujacy metode Create oraz BookAsync.")

    doc.add_heading("5.7 Cosmos DB - zapis i odczyt danych z kontenerow", level=2)
    doc.add_paragraph(
        "CosmosBookingRepository jest wspolna klasa do obslugi danych. "
        "Dzieki generykom ta sama logika moze dzialac dla klientow, fryzjerow, uslug, wizyt i uzytkownikow."
    )
    add_code(doc, """public async Task<T> CreateAsync(T entity, CancellationToken cancellationToken)
{
    entity.id = string.IsNullOrWhiteSpace(entity.id) ? Guid.NewGuid().ToString("N") : entity.id;
    entity.partitionKey = string.IsNullOrWhiteSpace(entity.partitionKey)
        ? entity.GetType().Name
        : entity.partitionKey;

    var container = containerResolver.Resolve<T>();
    var response = await container.CreateItemAsync(entity, new PartitionKey(entity.id), cancellationToken: cancellationToken);
    return response.Resource;
}

public async Task<T?> GetAsync(string id, CancellationToken cancellationToken)
{
    var container = containerResolver.Resolve<T>();
    var response = await container.ReadItemAsync<T>(id, new PartitionKey(id), cancellationToken: cancellationToken);
    return response.Resource;
}""", "Kod 7. Repository Cosmos DB - tworzenie i odczyt dokumentu.")
    doc.add_paragraph(
        "W Azure Data Explorer widac potem dokumenty JSON w kontenerach. Poniewaz w Twojej bazie partition key to /id, kod uzywa new PartitionKey(id)."
    )
    add_code(doc, """private Container ResolveContainer(Type entityType)
{
    if (entityType == typeof(Customer))
        return _database.GetContainer(_options.CustomersContainer);
    if (entityType == typeof(Hairdresser))
        return _database.GetContainer(_options.HairdressersContainer);
    if (entityType == typeof(Appointment))
        return _database.GetContainer(_options.AppointmentsContainer);
    if (entityType == typeof(SalonService))
        return _database.GetContainer(_options.ServicesContainer);
    if (entityType == typeof(AppUser))
        return _database.GetContainer(_options.UsersContainer);

    throw new InvalidOperationException($"No Cosmos DB container configured for {entityType.Name}.");
}""", "Kod 8. CosmosContainerResolver - podzial modeli na kontenery Cosmos DB.")
    add_placeholder(doc, "Zrzut 11. Visual Studio - CosmosBookingRepository.", "Wklej ekran kodu repozytorium oraz obok Azure Data Explorer z dokumentem w Cosmos DB.")

    doc.add_heading("5.8 Blob Storage - zdjecia salonu i fryzjerow", level=2)
    doc.add_paragraph(
        "Zdjecia nie sa trzymane w Cosmos DB. Do plikow binarnych uzywany jest Azure Blob Storage, a w bazie zapisywana jest tylko nazwa i adres pliku."
    )
    add_code(doc, """public async Task<UploadedPhoto> UploadAsync(string containerName, IFormFile file, CancellationToken cancellationToken)
{
    var container = new BlobContainerClient(connectionString, containerName);
    await container.CreateIfNotExistsAsync(cancellationToken: cancellationToken);

    var blobName = $"{Guid.NewGuid():N}{Path.GetExtension(file.FileName)}";
    var blob = container.GetBlobClient(blobName);

    await using var stream = file.OpenReadStream();
    await blob.UploadAsync(stream, new BlobHttpHeaders { ContentType = file.ContentType }, cancellationToken: cancellationToken);

    return new UploadedPhoto(file.FileName, blobName, blob.Uri.ToString());
}""", "Kod 9. BlobPhotoStorageService - wyslanie zdjecia do Blob Storage.")
    doc.add_paragraph(
        "Ten kod jest uruchamiany np. z kontrolera fryzjerow lub zdjec salonu. Po uploadzie aplikacja moze zwrocic URL zdjecia, a Azure przechowuje sam plik."
    )

    doc.add_heading("5.9 Queue Storage i Azure Functions - email po rezerwacji", level=2)
    doc.add_paragraph(
        "Po zapisaniu wizyty aplikacja nie wysyla maila bezposrednio w kontrolerze. "
        "Zamiast tego wysyla komunikat do kolejki. To pokazuje asynchroniczne przetwarzanie w chmurze."
    )
    add_code(doc, """var queue = new QueueClient(
    options.Value.Storage.ConnectionString,
    options.Value.Storage.AppointmentQueueName,
    new QueueClientOptions { MessageEncoding = QueueMessageEncoding.Base64 });

await queue.CreateIfNotExistsAsync(cancellationToken: cancellationToken);

var payload = JsonSerializer.Serialize(new
{
    appointmentId = appointment.id,
    customerId = appointment.CustomerId,
    customerEmail = customer?.Email,
    startAt = appointment.StartAt,
    endAt = appointment.EndAt
});

await queue.SendMessageAsync(payload, cancellationToken);""", "Kod 10. AzureQueueAppointmentBookedHandler - wyslanie komunikatu do booking-notifications.")
    add_code(doc, """[Function(nameof(AppointmentBookedQueueFunction))]
public async Task Run(
    [QueueTrigger("booking-notifications", Connection = "AzureWebJobsStorage")] string message,
    CancellationToken cancellationToken)
{
    logger.LogInformation("Queue message received by AppointmentBookedQueueFunction.");

    var appointment = JsonSerializer.Deserialize<AppointmentBookedMessage>(message, JsonOptions);
    await emailSender.SendAppointmentConfirmationAsync(appointment, cancellationToken);
}""", "Kod 11. AppointmentBookedQueueFunction - automatyczne uruchomienie po komunikacie w kolejce.")
    add_code(doc, """var client = new EmailClient(options.Value.ConnectionString);
var emailMessage = new EmailMessage(
    senderAddress: options.Value.SenderAddress,
    recipientAddress: message.CustomerEmail,
    content: new EmailContent("Potwierdzenie rezerwacji wizyty")
    {
        PlainText = BuildBody(message)
    });

await client.SendAsync(WaitUntil.Completed, emailMessage, cancellationToken);""", "Kod 12. AzureCommunicationEmailSender - wyslanie maila z ACS Email.")
    doc.add_paragraph(
        "Najwazniejsza zaleta tego rozwiazania: klient dostaje szybka odpowiedz z API, a email jest obslugiwany w tle przez Azure Function. "
        "Jesli wysylka maila chwilowo sie nie uda, problem dotyczy funkcji i kolejki, a nie samego tworzenia wizyty."
    )
    add_placeholder(doc, "Zrzut 12. Visual Studio - queue handler i Azure Function.", "Wklej ekran AzureQueueAppointmentBookedHandler, AppointmentBookedQueueFunction oraz Log Stream z przetwarzaniem kolejki.")

    doc.add_heading("5.10 SignalR i Easy Auth - komunikacja realtime oraz uzytkownik z chmury", level=2)
    doc.add_paragraph(
        "SignalR nie zastepuje kolejki ani maila. Jego rola jest inna: pokazuje natychmiastowe powiadomienie w otwartej aplikacji lub stronie testowej."
    )
    add_code(doc, """public async Task HandleAsync(Appointment appointment, CancellationToken cancellationToken)
{
    var payload = new
    {
        appointmentId = appointment.id,
        customerId = appointment.CustomerId,
        hairdresserId = appointment.HairdresserId,
        startAt = appointment.StartAt,
        endAt = appointment.EndAt,
        status = appointment.Status
    };

    await hubContext.Clients.All.SendAsync("appointmentBooked", payload, cancellationToken);
    await hubContext.Clients.Group(appointment.HairdresserId)
        .SendAsync("hairdresserAppointmentBooked", payload, cancellationToken);
}""", "Kod 13. SignalRAppointmentBookedHandler - powiadomienie realtime.")
    add_code(doc, """return new CurrentUserInfo(
    Provider: context.Request.Headers["X-MS-CLIENT-PRINCIPAL-IDP"].ToString(),
    ProviderUserId: context.Request.Headers["X-MS-CLIENT-PRINCIPAL-ID"].ToString(),
    DisplayName: context.Request.Headers["X-MS-CLIENT-PRINCIPAL-NAME"].ToString(),
    Email: context.Request.Headers["X-MS-CLIENT-PRINCIPAL-NAME"].ToString());""", "Kod 14. EasyAuthCurrentUserService - odczyt danych uzytkownika przekazanych przez Azure Easy Auth.")
    doc.add_paragraph(
        "Easy Auth dziala przed aplikacja ASP.NET Core. Po zalogowaniu przez Google lub GitHub Azure przekazuje do backendu naglowki X-MS-CLIENT-..., z ktorych aplikacja tworzy lub aktualizuje konto uzytkownika."
    )
    add_placeholder(doc, "Zrzut 13. SignalR i Easy Auth w dzialaniu.", "Wklej ekran strony signalr-test oraz ekran Authentication w App Service z providerami Google/GitHub.")

    doc.add_heading("6. Podsumowanie", level=1)
    doc.add_paragraph(
        "Projekt pokazuje praktyczne wykorzystanie chmury Azure w aplikacji backendowej. Najwazniejszy przeplyw laczy App Service, Cosmos DB, Queue Storage, Azure Functions i Azure Communication Services Email. "
        "Dodatkowo aplikacja korzysta z Blob Storage, SignalR, Easy Auth, Key Vault i monitoringu."
    )

    path = ROOT / "Raport_Chmura_Azure_HairSalonBooking_v2.docx"
    try:
        doc.save(path)
    except PermissionError:
        path = ROOT / "Raport_Chmura_Azure_HairSalonBooking_poprawiony.docx"
        doc.save(path)
    return path


def build_patterns_report() -> Path:
    doc = Document()
    style_doc(doc)
    add_header_footer(doc, "Raport wzorcow projektowych")
    add_identification(doc, "Raport 2: Wzorce projektowe")

    doc.add_heading("1. Tytul projektu", level=1)
    doc.add_paragraph("HairSalon Booking - backend aplikacji salonu fryzjerskiego do umawiania wizyt.")

    doc.add_heading("2. Cel projektu", level=1)
    doc.add_paragraph(
        "Celem projektu bylo stworzenie backendu aplikacji salonu fryzjerskiego oraz zastosowanie wzorcow projektowych w praktycznym kodzie. "
        "Wzorce mialy uporzadkowac logike biznesowa, dostep do danych, reakcje po utworzeniu wizyty oraz integracje z uslugami Azure."
    )

    doc.add_heading("3. Zalozenia projektu", level=1)
    doc.add_heading("3.1 Architektura pod wzorce", level=2)
    doc.add_paragraph(
        "Architektura aplikacji zostala podzielona na warstwy: API, Core oraz Functions. Dzieki temu kontrolery nie zawieraja calej logiki, a klasy maja wyrazne odpowiedzialnosci."
    )
    add_table(doc, ["Warstwa", "Odpowiedzialnosc", "Wzorce"], [
        ["API", "Kontrolery, konfiguracja, adaptery do Azure", "Dependency Injection, Adapter, Repository"],
        ["Core", "Modele, logika biznesowa, zdarzenia", "Facade, Factory Method, Observer/Event, Iterator"],
        ["Functions", "Przetwarzanie w tle i email", "Adapter, Options Pattern, event-driven processing"],
    ])

    doc.add_heading("3.2 Wybrane technologie z uzasadnieniem", level=2)
    add_table(doc, ["Technologia", "Znaczenie dla wzorcow"], [
        ["C# / .NET 8", "Wspiera interfejsy, generyki, async/await, dependency injection i zdarzenia."],
        ["ASP.NET Core", "Wbudowany kontener DI pozwala rejestrowac interfejsy i implementacje."],
        ["Azure Cosmos DB", "Wymaga oddzielenia dostepu do danych przez Repository."],
        ["Azure Queue + Functions", "Naturalnie pasuje do Observer/Event i Handlerow po utworzeniu wizyty."],
        ["Azure SDK", "Zostalo ukryte za Adapterami, aby nie uzalezniac kontrolerow od klas SDK."],
    ])

    doc.add_heading("4. Funkcjonalnosc a wzorce", level=1)
    add_table(doc, ["Funkcjonalnosc", "Uzyte wzorce", "Opis"], [
        ["Tworzenie wizyty", "Facade, Repository, Observer/Event", "Facade koordynuje proces rezerwacji, Repository zapisuje dane, Event uruchamia reakcje."],
        ["CRUD klientow/fryzjerow/uslug", "Repository, Dependency Injection", "Kontrolery korzystaja z IBookingRepository zamiast bezposrednio z Cosmos DB."],
        ["Email po rezerwacji", "Observer/Event, Handler, Adapter", "Po zdarzeniu wizyta trafia do kolejki, a Function uzywa IEmailSender."],
        ["Powiadomienia SignalR", "Handler, Adapter", "Osobny handler wysyla powiadomienie realtime."],
        ["Zdjecia salonu/fryzjerow", "Adapter, Options Pattern", "IPhotoStorageService ukrywa Blob Storage."],
        ["Wolne terminy", "Iterator", "AvailabilitySlotCollection udostepnia sloty jako IEnumerable."],
    ])
    add_placeholder(doc, "Zrzut 1. Solution Explorer z warstwami projektu.", "Wklej ekran Visual Studio pokazujacy projekty Api, Core i Functions.")
    add_placeholder(doc, "Zrzut 2. Swagger - funkcjonalnosci CRUD.", "Wklej ekran Swaggera potwierdzajacy funkcjonalnosci aplikacji.")

    doc.add_heading("5. Przygotowanie projektu - wzorce w kodzie", level=1)
    doc.add_heading("5.1 Facade", level=2)
    doc.add_paragraph(
        "Facade znajduje sie w klasie AppointmentBookingFacade. Wzorzec ukrywa zlozony proces rezerwacji wizyty za jedna metoda BookAsync."
    )
    add_code(doc, """public async Task<Appointment> BookAsync(Appointment appointment, CancellationToken cancellationToken)
{
    var customer = await customers.GetAsync(appointment.CustomerId, cancellationToken);
    var hairdresser = await hairdressers.GetAsync(appointment.HairdresserId, cancellationToken);
    var service = await services.GetAsync(appointment.SalonServiceId, cancellationToken);
    appointment.EndAt = appointment.StartAt.AddMinutes(service.DurationMinutes);
    await EnsureSlotIsFreeAsync(appointment, cancellationToken);
    var created = await appointments.CreateAsync(appointment, cancellationToken);
    await eventPublisher.PublishAppointmentBookedAsync(created, cancellationToken);
    return created;
}""", "Kod 1. Fragment wzorca Facade w AppointmentBookingFacade.")
    add_placeholder(doc, "Zrzut 3. Kod AppointmentBookingFacade.", "Wklej ekran z Visual Studio pokazujacy metode BookAsync.")

    doc.add_heading("5.2 Repository", level=2)
    doc.add_paragraph(
        "Repository oddziela kontrolery i logike biznesowa od szczegolow Cosmos DB. Aplikacja uzywa interfejsu IBookingRepository<T>."
    )
    add_code(doc, """public interface IBookingRepository<T> where T : BookingEntity
{
    Task<IReadOnlyList<T>> GetAllAsync(CancellationToken cancellationToken);
    Task<T?> GetAsync(string id, CancellationToken cancellationToken);
    Task<T> CreateAsync(T entity, CancellationToken cancellationToken);
    Task<T> UpsertAsync(T entity, CancellationToken cancellationToken);
    Task DeleteAsync(string id, CancellationToken cancellationToken);
}""", "Kod 2. Interfejs Repository.")
    add_code(doc, """public sealed class CosmosBookingRepository<T>(ICosmosContainerResolver containerResolver)
    : IBookingRepository<T> where T : BookingEntity, new()
{
    private readonly Container _container = containerResolver.GetContainer<T>();
}""", "Kod 3. Implementacja Repository dla Cosmos DB.")

    doc.add_heading("5.3 Observer/Event oraz Strategy/Handler", level=2)
    doc.add_paragraph(
        "Po utworzeniu wizyty publikowane jest zdarzenie. Kilka handlerow reaguje na nie niezaleznie: Blob, Queue i SignalR."
    )
    add_code(doc, """foreach (var handler in handlers)
{
    await handler.HandleAsync(appointment, cancellationToken);
}""", "Kod 4. Publisher uruchamia wszystkie handlery zdarzenia.")
    add_code(doc, """public interface IAppointmentBookedHandler
{
    Task HandleAsync(Appointment appointment, CancellationToken cancellationToken);
}""", "Kod 5. Wspolny interfejs handlerow.")
    add_table(doc, ["Handler", "Zadanie"], [
        ["AzureBlobAppointmentSummaryHandler", "Zapisuje podsumowanie wizyty do Blob Storage."],
        ["AzureQueueAppointmentBookedHandler", "Wrzuca komunikat do kolejki booking-notifications."],
        ["SignalRAppointmentBookedHandler", "Wysyla powiadomienie realtime appointmentBooked."],
    ])
    add_placeholder(doc, "Zrzut 4. Kod BookingEventPublisher lub IAppointmentBookedHandler.", "Wklej ekran kodu pokazujacy Observer/Event albo Strategy/Handler.")

    doc.add_heading("5.4 Factory Method", level=2)
    doc.add_paragraph("Factory Method centralizuje tworzenie encji domenowych w klasie BookingEntityFactory.")
    add_code(doc, """public BookingEntity Create(EntityKind kind) => kind switch
{
    EntityKind.Customer => new Customer(),
    EntityKind.Hairdresser => new Hairdresser(),
    EntityKind.SalonService => new SalonService(),
    EntityKind.Appointment => new Appointment(),
    EntityKind.SalonPhoto => new SalonPhoto(),
    EntityKind.AppUser => new AppUser(),
    _ => throw new ArgumentOutOfRangeException(...)
};""", "Kod 6. Factory Method dla encji domenowych.")

    doc.add_heading("5.5 Dependency Injection, Singleton i Adapter", level=2)
    doc.add_paragraph(
        "Dependency Injection laczy wszystkie wzorce. Singleton zostal uzyty m.in. dla CosmosClient, a Adaptery ukrywaja uslugi Azure."
    )
    add_code(doc, """services.AddScoped<IAppointmentBookingFacade, AppointmentBookingFacade>();
services.AddScoped<IBookingEventPublisher, BookingEventPublisher>();
services.AddScoped(typeof(IBookingRepository<>), typeof(CosmosBookingRepository<>));
services.AddSingleton(_ => new CosmosClient(options.Cosmos.ConnectionString));
services.AddScoped<IPhotoStorageService, BlobPhotoStorageService>();""", "Kod 7. Rejestracja zaleznosci w ServiceCollectionExtensions.")
    add_table(doc, ["Adapter", "Ukrywana usluga"], [
        ["BlobPhotoStorageService", "Azure Blob Storage"],
        ["EasyAuthCurrentUserService", "Naglowki Azure Easy Auth"],
        ["AzureCommunicationEmailSender", "Azure Communication Services Email"],
    ])

    doc.add_heading("5.6 Iterator i Options Pattern", level=2)
    doc.add_paragraph(
        "Iterator pojawia sie w AvailabilitySlotCollection, ktora generuje sloty wizyt i udostepnia je przez IEnumerable. Options Pattern porzadkuje konfiguracje Azure."
    )
    add_code(doc, """public sealed class AvailabilitySlotCollection : IEnumerable<DateTimeOffset>
{
    public IEnumerator<DateTimeOffset> GetEnumerator() => _slots.GetEnumerator();
}""", "Kod 8. Iterator dla slotow wizyt.")
    add_code(doc, """services.Configure<AzureBookingOptions>(configuration.GetSection("Azure"));
services.Configure<EmailOptions>(context.Configuration.GetSection("Email"));""", "Kod 9. Options Pattern dla konfiguracji.")

    doc.add_heading("6. Podsumowanie", level=1)
    doc.add_paragraph(
        "W projekcie zastosowano wzorce w miejscach wynikajacych z architektury aplikacji. Najwazniejsze to Facade, Repository, Observer/Event, Strategy/Handler i Factory Method. "
        "Pozostale wzorce wspieraja integracje z Azure, konfiguracje i utrzymanie kodu."
    )
    add_note(doc, "Przed oddaniem raportu warto wkleic zrzuty ekranow z Visual Studio pokazujace konkretne klasy wzorcow oraz screen ze Swaggera potwierdzajacy funkcjonalnosc.")

    path = ROOT / "Raport_Wzorce_Projektowe_HairSalonBooking_v2.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_cloud_report())
    print(build_patterns_report())
