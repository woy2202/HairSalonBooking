from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
AZURE_IMAGE = Path(r"C:\Users\woy22\Downloads\azure.jpg")


BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
GREEN = "E2F0D9"
ORANGE = "FCE4D6"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    p.paragraph_format.space_after = Pt(2)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.08

    for name, size, color in [
        ("Title", 24, BLUE),
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 11.5, BLUE),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def add_cover(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(80, 80, 80)

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Projekt", "HairSalon Booking - backend aplikacji salonu fryzjerskiego"),
        ("Technologie", ".NET 8, C#, ASP.NET Core, Azure"),
        ("Autor", "Wiktor Gochnio"),
        ("Data", "18.05.2026"),
    ]
    for row, (k, v) in zip(table.rows, rows):
        set_cell_text(row.cells[0], k, True)
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[1], v)
    doc.add_page_break()


def add_info_box(doc: Document, title: str, body: str, fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p.add_run("\n" + body)
    doc.add_paragraph()


def add_screenshot_placeholder(doc: Document, caption: str, description: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("[MIEJSCE NA ZRZUT EKRANU]\n")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p.add_run(description)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(16)
        paragraph.paragraph_format.space_after = Pt(16)
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)
    doc.add_paragraph()


def add_code_block(doc: Document, text: str, caption: str | None = None) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8F8F8")
    cell.text = ""
    p = cell.paragraphs[0]
    for line_no, line in enumerate(text.splitlines()):
        if line_no:
            p.add_run("\n")
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
    if caption:
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(9)
    doc.add_paragraph()


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, True)
        set_cell_shading(table.rows[0].cells[i], LIGHT_BLUE)
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            set_cell_text(cells[i], value)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_cloud_report() -> Path:
    doc = Document()
    style_document(doc)
    add_cover(
        doc,
        "Raport projektu - chmura Azure",
        "Backend aplikacji salonu fryzjerskiego do rezerwacji wizyt",
    )

    doc.add_heading("1. Cel projektu", level=1)
    doc.add_paragraph(
        "Celem projektu bylo przygotowanie backendu aplikacji dla salonu fryzjerskiego, ktory pozwala zarzadzac klientami, fryzjerami, uslugami oraz wizytami. "
        "Aplikacja zostala uruchomiona w chmurze Microsoft Azure i wykorzystuje kilka uslug chmurowych, aby pokazac praktyczne uzycie architektury rozproszonej."
    )
    add_info_box(
        doc,
        "Najwazniejszy rezultat",
        "Po utworzeniu wizyty w API aplikacja zapisuje dane w Cosmos DB, wysyla powiadomienie realtime przez SignalR, umieszcza komunikat w kolejce, a Azure Function wysyla klientowi email z potwierdzeniem wizyty.",
        GREEN,
    )

    doc.add_heading("2. Zalozenia projektu", level=1)
    doc.add_heading("2.1 Architektura", level=2)
    doc.add_paragraph(
        "Projekt zostal podzielony na trzy glowne projekty: warstwe API, warstwe Core z logika domenowa oraz projekt Azure Functions odpowiedzialny za przetwarzanie w tle."
    )
    add_table(
        doc,
        ["Projekt", "Rola"],
        [
            ["HairSalon.Booking.Api", "Aplikacja ASP.NET Core Web API. Udostepnia kontrolery REST, Swagger, SignalR Hub i laczy sie z Cosmos DB, Blob Storage oraz Queue Storage."],
            ["HairSalon.Booking.Core", "Warstwa domenowa. Zawiera modele, interfejsy repozytoriow, wzorce projektowe, logike rezerwacji i zdarzenia."],
            ["HairSalon.Booking.Functions", "Azure Functions. Odbiera komunikaty z kolejki i wysyla email z potwierdzeniem rezerwacji."],
        ],
        [4.5, 11.5],
    )
    add_code_block(
        doc,
        "Klient / Swagger / Frontend\n"
        "        |\n"
        "        v\n"
        "Azure App Service: booking-api\n"
        "        |\n"
        "        +--> Azure Cosmos DB - dane aplikacji\n"
        "        +--> Azure Blob Storage - zdjecia i podsumowania\n"
        "        +--> Azure Queue Storage - komunikaty o wizytach\n"
        "        +--> Azure SignalR Service - powiadomienia realtime\n"
        "        |\n"
        "        v\n"
        "Azure Function: booking-function\n"
        "        |\n"
        "        v\n"
        "Azure Communication Services Email - email do klienta",
        "Rysunek 1. Schemat logiczny architektury aplikacji.",
    )

    doc.add_heading("2.2 Wybrane technologie i uzasadnienie", level=2)
    add_table(
        doc,
        ["Technologia / usluga", "Zastosowanie", "Uzasadnienie"],
        [
            [".NET 8 / C#", "Backend REST API", "Technologia zgodna z wymaganiami projektu, dobra integracja z Azure i Visual Studio."],
            ["ASP.NET Core Web API", "Kontrolery CRUD i Swagger", "Pozwala szybko zbudowac testowalne API dla klientow, fryzjerow, wizyt i uslug."],
            ["Azure App Service", "Hosting API", "Latwe wdrozenie aplikacji webowej bez zarzadzania serwerem."],
            ["Azure Cosmos DB", "Baza danych NoSQL", "Dobre dopasowanie do dokumentow JSON i kontenerow: customers, barbers, appointments, services, users."],
            ["Azure Blob Storage", "Zdjecia salonu/fryzjerow i pliki tekstowe z podsumowaniem", "Tanie i proste przechowywanie plikow binarnych."],
            ["Azure Queue Storage", "Kolejka zdarzen wizyt", "Oddziela rezerwacje wizyty od wysylki maila."],
            ["Azure Functions", "Przetwarzanie w tle", "Serverless: funkcja uruchamia sie po komunikacie w kolejce."],
            ["Azure Communication Services Email", "Potwierdzenia email", "Profesjonalna usluga do wysylki wiadomosci z chmury."],
            ["Azure SignalR Service", "Powiadomienia realtime", "Pozwala wyslac natychmiastowa informacje o nowej wizycie."],
            ["Azure Key Vault", "Sekrety i connection stringi", "Bezpieczniejsze przechowywanie kluczy poza kodem."],
            ["Application Insights / Monitor", "Logi i diagnostyka", "Pomaga sprawdzac requesty, bledy i dzialanie Function App."],
            ["Easy Auth", "Logowanie Google/GitHub", "Szybkie zabezpieczenie API przez Azure App Service Authentication."],
        ],
        [4.2, 5.3, 6.5],
    )

    doc.add_heading("3. Funkcjonalnosc", level=1)
    add_bullets(
        doc,
        [
            "Pelny CRUD klientow: dodawanie, odczyt, aktualizacja i usuwanie danych klienta.",
            "Pelny CRUD fryzjerow: zarzadzanie profilami, specjalizacja, aktywnosc i zdjecia.",
            "Pelny CRUD uslug salonu: nazwa, opis, cena, czas trwania i dostepnosc.",
            "Pelny CRUD wizyt: rezerwacja, odczyt, aktualizacja i usuwanie terminow.",
            "Walidacja rezerwacji: sprawdzenie klienta, fryzjera, uslugi i kolizji terminow.",
            "Automatyczne wyslanie maila po utworzeniu wizyty.",
            "Powiadomienie realtime przez SignalR po utworzeniu wizyty.",
            "Logowanie przez Google/GitHub z wykorzystaniem Easy Auth.",
            "Role uzytkownikow: klient, fryzjer, admin.",
            "Mozliwosc dodawania zdjec salonu i fryzjerow do Blob Storage.",
        ],
    )

    doc.add_heading("3.1 Swagger i testowanie CRUD", level=2)
    add_screenshot_placeholder(
        doc,
        "Zrzut 1. Swagger UI z widocznymi kontrolerami aplikacji.",
        "Wstaw zrzut ekranu ze Swaggera, na ktorym widac kontrolery Customers, Hairdressers, Appointments, SalonServices, Auth/Admin.",
    )
    add_screenshot_placeholder(
        doc,
        "Zrzut 2. Przyklad utworzenia wizyty przez POST /api/Appointments.",
        "Wstaw zrzut ekranu z odpowiedzia 201 Created po dodaniu wizyty.",
    )

    doc.add_heading("3.2 Wysylka maila po rezerwacji", level=2)
    doc.add_paragraph(
        "Po utworzeniu wizyty API wrzuca komunikat do kolejki booking-notifications. Azure Function odbiera komunikat i wysyla email z potwierdzeniem wizyty przez Azure Communication Services Email."
    )
    add_screenshot_placeholder(
        doc,
        "Zrzut 3. Wiadomosc email z potwierdzeniem wizyty.",
        "Wstaw zrzut ekranu otrzymanego maila od DoNotReply@...azurecomm.net.",
    )
    add_screenshot_placeholder(
        doc,
        "Zrzut 4. Log stream Azure Function z obsluga komunikatu z kolejki.",
        "Wstaw zrzut ekranu logow Function App pokazujacy odebranie komunikatu i wysylke emaila.",
    )

    doc.add_heading("3.3 SignalR", level=2)
    doc.add_paragraph(
        "SignalR sluzy do natychmiastowych powiadomien online. Po utworzeniu wizyty aplikacja wysyla komunikat appointmentBooked do podlaczonych klientow."
    )
    add_screenshot_placeholder(
        doc,
        "Zrzut 5. Test SignalR z komunikatem appointmentBooked.",
        "Wstaw zrzut ekranu strony /api/Notifications/signalr-test z komunikatem Connected to SignalR.",
    )

    doc.add_heading("4. Przygotowanie projektu", level=1)
    doc.add_heading("4.1 Utworzenie zasobow Azure", level=2)
    if AZURE_IMAGE.exists():
        doc.add_paragraph("Ponizszy zrzut pokazuje grupe zasobow Booking w Azure Portal.")
        doc.add_picture(str(AZURE_IMAGE), width=Cm(16))
        cap = doc.add_paragraph("Zrzut 6. Grupa zasobow Booking w Azure Portal.")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
    else:
        add_screenshot_placeholder(
            doc,
            "Zrzut 6. Grupa zasobow Booking w Azure Portal.",
            "Wstaw zrzut ekranu z Azure Portal pokazujacy App Service, Cosmos DB, Function App, Storage Account i SignalR.",
        )

    doc.add_heading("4.2 Konfiguracja Cosmos DB", level=2)
    doc.add_paragraph("W Cosmos DB utworzono baze BookingApkaDB oraz osobne kontenery dla danych aplikacji.")
    add_table(
        doc,
        ["Kontener", "Przechowywane dane", "Partition key"],
        [
            ["customers", "Klienci salonu", "/id"],
            ["barbers", "Fryzjerzy", "/id"],
            ["services", "Uslugi salonu", "/id"],
            ["appointments", "Wizyty", "/id"],
            ["users", "Konta aplikacji i role", "/id"],
            ["salon-photos", "Zdjecia salonu", "/id"],
        ],
        [4.5, 8.0, 3.0],
    )
    add_screenshot_placeholder(
        doc,
        "Zrzut 7. Cosmos DB Data Explorer z kontenerami projektu.",
        "Wstaw zrzut ekranu z Data Explorer pokazujacy kontenery customers, barbers, appointments, services, users.",
    )

    doc.add_heading("4.3 Konfiguracja App Service i Function App", level=2)
    add_table(
        doc,
        ["Zasob", "Najwazniejsze ustawienia"],
        [
            ["booking-api", "Runtime .NET 8, App Service Authentication, Application settings dla Cosmos, Storage, SignalR, Key Vault."],
            ["booking-function", "Runtime .NET 8 isolated, AzureWebJobsStorage, Email__ConnectionString, Email__SenderAddress."],
            ["bookingfryzjer", "Storage Account z Queue Storage i Blob Storage."],
        ],
        [4.5, 11.5],
    )

    doc.add_heading("4.4 Kluczowe fragmenty kodu", level=2)
    add_code_block(
        doc,
        "app.MapControllers();\n"
        "app.MapHub<BookingNotificationsHub>(\"/hubs/booking-notifications\");",
        "Kod 1. Mapowanie kontrolerow i huba SignalR w Program.cs.",
    )
    add_code_block(
        doc,
        "[QueueTrigger(\"booking-notifications\", Connection = \"AzureWebJobsStorage\")]\n"
        "public async Task Run(string message, CancellationToken cancellationToken)\n"
        "{\n"
        "    var appointment = JsonSerializer.Deserialize<AppointmentBookedMessage>(message, JsonOptions);\n"
        "    await emailSender.SendAppointmentConfirmationAsync(appointment, cancellationToken);\n"
        "}",
        "Kod 2. Azure Function reagujaca na komunikaty z kolejki.",
    )
    add_code_block(
        doc,
        "await queue.SendMessageAsync(payload, cancellationToken);",
        "Kod 3. Wyslanie komunikatu o wizycie do Azure Queue Storage.",
    )

    doc.add_heading("5. Podsumowanie", level=1)
    doc.add_paragraph(
        "Projekt pokazuje kompletna integracje backendu C# z uslugami Azure. Aplikacja nie ogranicza sie do CRUD, ale wykorzystuje elementy chmurowe: baze NoSQL, kolejki, funkcje serverless, powiadomienia realtime, przechowywanie plikow, monitoring i wysylke emaili."
    )
    add_info_box(
        doc,
        "Wniosek",
        "Architektura jest modularna: API odpowiada za logike i endpointy, Cosmos DB za dane, Storage za pliki i kolejki, Function App za zadania w tle, a SignalR za realtime.",
        GREEN,
    )

    path = ROOT / "Raport_1_Chmura_Azure_HairSalonBooking.docx"
    doc.save(path)
    return path


def build_patterns_report() -> Path:
    doc = Document()
    style_document(doc)
    add_cover(
        doc,
        "Raport projektu - wzorce projektowe",
        "Wzorce zastosowane w backendzie aplikacji salonu fryzjerskiego",
    )

    doc.add_heading("1. Cel projektu", level=1)
    doc.add_paragraph(
        "Celem tej czesci raportu jest przedstawienie wzorcow projektowych uzytych w aplikacji HairSalon Booking. Raport opisuje, gdzie wzorce wystepuja w kodzie, dlaczego zostaly uzyte oraz jaki problem rozwiazuja."
    )
    add_info_box(
        doc,
        "Najwazniejsza mysl",
        "Wzorce zostaly uzyte po to, aby kontrolery byly proste, logika biznesowa byla oddzielona od infrastruktury, a reakcje po utworzeniu wizyty mozna bylo rozwijac bez przepisywania calej aplikacji.",
        GREEN,
    )

    doc.add_heading("2. Zalozenia architektoniczne", level=1)
    doc.add_paragraph(
        "Projekt zostal podzielony na warstwy, co naturalnie wspiera zastosowanie wzorcow. Warstwa API obsluguje HTTP, warstwa Core zawiera modele i logike biznesowa, a projekt Functions realizuje przetwarzanie w tle."
    )
    add_table(
        doc,
        ["Warstwa", "Odpowiedzialnosc", "Przykladowe wzorce"],
        [
            ["API", "Kontrolery, konfiguracja, integracje Azure", "Dependency Injection, Adapter, Repository"],
            ["Core", "Logika domenowa, modele, zdarzenia", "Facade, Factory Method, Observer, Iterator"],
            ["Functions", "Obsluga kolejki i maili", "Adapter, Options Pattern, event-driven processing"],
        ],
        [3.5, 6.5, 6.0],
    )

    doc.add_heading("3. Lista uzytych wzorcow", level=1)
    add_table(
        doc,
        ["Wzorzec", "Gdzie wystepuje", "Po co zostal uzyty"],
        [
            ["Facade", "AppointmentBookingFacade", "Ukrywa zlozony proces rezerwacji wizyty za jedna metoda BookAsync."],
            ["Repository", "IBookingRepository, CosmosBookingRepository", "Oddziela aplikacje od szczegolow Cosmos DB."],
            ["Observer/Event", "BookingEventPublisher", "Pozwala wielu komponentom reagowac na utworzenie wizyty."],
            ["Strategy/Handler", "IAppointmentBookedHandler i implementacje", "Kazda reakcja po wizycie ma ten sam interfejs, ale inne zachowanie."],
            ["Factory Method", "BookingEntityFactory", "Centralizuje tworzenie encji domenowych."],
            ["Dependency Injection", "ServiceCollectionExtensions", "Wstrzykuje zaleznosci przez interfejsy."],
            ["Adapter/Wrapper", "BlobPhotoStorageService, EasyAuthCurrentUserService, AzureCommunicationEmailSender", "Ukrywa szczegoly uslug Azure za prostymi interfejsami."],
            ["Singleton", "CosmosClient, CosmosContainerResolver", "Wspoldzieli kosztowne obiekty infrastrukturalne."],
            ["Iterator", "AvailabilitySlotCollection", "Udostepnia sloty wizyt jako kolekcje IEnumerable."],
            ["Options Pattern", "AzureBookingOptions, EmailOptions", "Porzadkuje konfiguracje Azure."],
        ],
        [3.5, 5.5, 7.0],
    )

    doc.add_heading("4. Facade", level=1)
    doc.add_paragraph(
        "Facade jest widoczny w klasie AppointmentBookingFacade. Klasa koordynuje caly proces rezerwacji wizyty i ukrywa jego zlozonosc przed kontrolerem."
    )
    add_code_block(
        doc,
        "public sealed class AppointmentBookingFacade(\n"
        "    IBookingRepository<Customer> customers,\n"
        "    IBookingRepository<Hairdresser> hairdressers,\n"
        "    IBookingRepository<SalonService> services,\n"
        "    IBookingRepository<Appointment> appointments,\n"
        "    IBookingEventPublisher eventPublisher) : IAppointmentBookingFacade\n"
        "{\n"
        "    public async Task<Appointment> BookAsync(Appointment appointment, CancellationToken cancellationToken)\n"
        "    {\n"
        "        var customer = await customers.GetAsync(appointment.CustomerId, cancellationToken);\n"
        "        var hairdresser = await hairdressers.GetAsync(appointment.HairdresserId, cancellationToken);\n"
        "        var service = await services.GetAsync(appointment.SalonServiceId, cancellationToken);\n"
        "        appointment.EndAt = appointment.StartAt.AddMinutes(service.DurationMinutes);\n"
        "        await EnsureSlotIsFreeAsync(appointment, cancellationToken);\n"
        "        var created = await appointments.CreateAsync(appointment, cancellationToken);\n"
        "        await eventPublisher.PublishAppointmentBookedAsync(created, cancellationToken);\n"
        "        return created;\n"
        "    }\n"
        "}",
        "Kod 1. Facade procesu rezerwacji wizyty.",
    )
    doc.add_paragraph(
        "Kontroler AppointmentsController nie musi znac szczegolow walidacji ani reakcji po utworzeniu wizyty. Wywoluje tylko bookingFacade.BookAsync. To sprawia, ze kontroler jest krotki i czytelny."
    )

    doc.add_heading("5. Repository", level=1)
    doc.add_paragraph(
        "Repository oddziela logike aplikacji od sposobu przechowywania danych. Kontrolery i serwisy pracuja na interfejsie IBookingRepository, a nie bezposrednio na CosmosClient."
    )
    add_code_block(
        doc,
        "public interface IBookingRepository<T> where T : BookingEntity\n"
        "{\n"
        "    Task<IReadOnlyList<T>> GetAllAsync(CancellationToken cancellationToken);\n"
        "    Task<T?> GetAsync(string id, CancellationToken cancellationToken);\n"
        "    Task<T> CreateAsync(T entity, CancellationToken cancellationToken);\n"
        "    Task<T> UpsertAsync(T entity, CancellationToken cancellationToken);\n"
        "    Task DeleteAsync(string id, CancellationToken cancellationToken);\n"
        "}",
        "Kod 2. Interfejs Repository.",
    )
    add_code_block(
        doc,
        "public sealed class CosmosBookingRepository<T>(ICosmosContainerResolver containerResolver)\n"
        "    : IBookingRepository<T> where T : BookingEntity, new()\n"
        "{\n"
        "    private readonly Container _container = containerResolver.GetContainer<T>();\n"
        "    public async Task<T> CreateAsync(T entity, CancellationToken cancellationToken)\n"
        "    {\n"
        "        entity.id = string.IsNullOrWhiteSpace(entity.id) ? Guid.NewGuid().ToString(\"N\") : entity.id;\n"
        "        var response = await _container.CreateItemAsync(entity, new PartitionKey(entity.id), cancellationToken: cancellationToken);\n"
        "        return response.Resource;\n"
        "    }\n"
        "}",
        "Kod 3. Implementacja Repository dla Cosmos DB.",
    )
    doc.add_paragraph(
        "Dzieki temu ten sam interfejs moze miec rozne implementacje, np. CosmosBookingRepository dla Azure oraz InMemoryBookingRepository do pracy lokalnej."
    )

    doc.add_heading("6. Observer / Event", level=1)
    doc.add_paragraph(
        "Observer/Event zostal uzyty do reakcji po utworzeniu wizyty. Rezerwacja wizyty publikuje zdarzenie, a rozne komponenty moga na nie odpowiedziec."
    )
    add_code_block(
        doc,
        "public sealed class BookingEventPublisher(IEnumerable<IAppointmentBookedHandler> handlers) : IBookingEventPublisher\n"
        "{\n"
        "    public event EventHandler<AppointmentBookedEventArgs>? AppointmentBooked;\n"
        "\n"
        "    public async Task PublishAppointmentBookedAsync(Appointment appointment, CancellationToken cancellationToken)\n"
        "    {\n"
        "        AppointmentBooked?.Invoke(this, new AppointmentBookedEventArgs(appointment));\n"
        "        foreach (var handler in handlers)\n"
        "        {\n"
        "            await handler.HandleAsync(appointment, cancellationToken);\n"
        "        }\n"
        "    }\n"
        "}",
        "Kod 4. Publisher zdarzenia utworzenia wizyty.",
    )
    doc.add_paragraph(
        "Ten wzorzec pozwala dodawac nowe reakcje, np. SMS albo wpis do raportu, bez zmieniania kontrolera i glownej logiki rezerwacji."
    )

    doc.add_heading("7. Strategy / Handler", level=1)
    doc.add_paragraph(
        "Strategy/Handler jest widoczny w interfejsie IAppointmentBookedHandler. Kazda implementacja realizuje inna reakcje na to samo zdarzenie."
    )
    add_code_block(
        doc,
        "public interface IAppointmentBookedHandler\n"
        "{\n"
        "    Task HandleAsync(Appointment appointment, CancellationToken cancellationToken);\n"
        "}",
        "Kod 5. Wspolny interfejs handlerow.",
    )
    add_table(
        doc,
        ["Handler", "Zachowanie"],
        [
            ["AzureBlobAppointmentSummaryHandler", "Zapisuje tekstowe podsumowanie wizyty do Blob Storage."],
            ["AzureQueueAppointmentBookedHandler", "Wrzuca komunikat do Azure Queue Storage, aby Function wyslala email."],
            ["SignalRAppointmentBookedHandler", "Wysyla powiadomienie realtime przez SignalR."],
        ],
        [6.0, 10.0],
    )
    add_code_block(
        doc,
        "services.AddScoped<IAppointmentBookedHandler, AzureBlobAppointmentSummaryHandler>();\n"
        "services.AddScoped<IAppointmentBookedHandler, AzureQueueAppointmentBookedHandler>();\n"
        "services.AddScoped<IAppointmentBookedHandler, SignalRAppointmentBookedHandler>();",
        "Kod 6. Rejestracja wielu strategii/handlerow w Dependency Injection.",
    )

    doc.add_heading("8. Factory Method", level=1)
    doc.add_paragraph(
        "Factory Method wystepuje w BookingEntityFactory. Fabryka tworzy odpowiedni typ encji na podstawie EntityKind."
    )
    add_code_block(
        doc,
        "public BookingEntity Create(EntityKind kind) => kind switch\n"
        "{\n"
        "    EntityKind.Customer => new Customer(),\n"
        "    EntityKind.Hairdresser => new Hairdresser(),\n"
        "    EntityKind.SalonService => new SalonService(),\n"
        "    EntityKind.Appointment => new Appointment(),\n"
        "    EntityKind.SalonPhoto => new SalonPhoto(),\n"
        "    EntityKind.AppUser => new AppUser(),\n"
        "    _ => throw new ArgumentOutOfRangeException(...)\n"
        "};",
        "Kod 7. Factory Method dla encji domenowych.",
    )
    doc.add_paragraph(
        "Dzieki temu tworzenie obiektow domenowych jest scentralizowane i latwo rozszerzalne."
    )

    doc.add_heading("9. Dependency Injection i Singleton", level=1)
    doc.add_paragraph(
        "Dependency Injection spina wszystkie wzorce. Klasy otrzymuja zaleznosci przez konstruktory, a implementacje sa rejestrowane w ServiceCollectionExtensions."
    )
    add_code_block(
        doc,
        "services.AddScoped<IAppointmentBookingFacade, AppointmentBookingFacade>();\n"
        "services.AddScoped<IBookingEventPublisher, BookingEventPublisher>();\n"
        "services.AddScoped(typeof(IBookingRepository<>), typeof(CosmosBookingRepository<>));\n"
        "services.AddSingleton(_ => new CosmosClient(options.Cosmos.ConnectionString));\n"
        "services.AddSingleton<ICosmosContainerResolver, CosmosContainerResolver>();",
        "Kod 8. Dependency Injection i Singleton w konfiguracji aplikacji.",
    )
    doc.add_paragraph(
        "CosmosClient jest singletonem, poniewaz jest kosztownym obiektem SDK i powinien byc wspoldzielony zamiast tworzony przy kazdym requescie."
    )

    doc.add_heading("10. Adapter / Wrapper", level=1)
    doc.add_paragraph(
        "Adaptery ukrywaja szczegoly zewnetrznych uslug Azure za interfejsami aplikacji. Dzieki temu kontrolery i logika biznesowa nie musza znac bezposrednio Azure SDK."
    )
    add_table(
        doc,
        ["Adapter", "Interfejs", "Co ukrywa"],
        [
            ["BlobPhotoStorageService", "IPhotoStorageService", "Tworzenie BlobContainerClient, upload i usuwanie plikow."],
            ["EasyAuthCurrentUserService", "ICurrentUserService", "Naglowki Easy Auth X-MS-CLIENT-PRINCIPAL."],
            ["AzureCommunicationEmailSender", "IEmailSender", "EmailClient i szczegoly wysylania maila przez ACS."],
        ],
        [5.0, 4.5, 6.5],
    )
    add_code_block(
        doc,
        "public interface IPhotoStorageService\n"
        "{\n"
        "    Task<UploadedPhoto> UploadAsync(string containerName, IFormFile file, CancellationToken cancellationToken);\n"
        "    Task DeleteAsync(string containerName, string blobName, CancellationToken cancellationToken);\n"
        "}",
        "Kod 9. Interfejs ukrywajacy szczegoly Blob Storage.",
    )

    doc.add_heading("11. Iterator", level=1)
    doc.add_paragraph(
        "Iterator wystepuje w AvailabilitySlotCollection. Klasa generuje dostepne sloty wizyt i udostepnia je jako IEnumerable<DateTimeOffset>."
    )
    add_code_block(
        doc,
        "public sealed class AvailabilitySlotCollection : IEnumerable<DateTimeOffset>\n"
        "{\n"
        "    private readonly List<DateTimeOffset> _slots = [];\n"
        "    public IEnumerator<DateTimeOffset> GetEnumerator() => _slots.GetEnumerator();\n"
        "    IEnumerator IEnumerable.GetEnumerator() => GetEnumerator();\n"
        "}",
        "Kod 10. Iterator dla slotow wizyt.",
    )
    doc.add_paragraph(
        "Dzieki temu AppointmentBookingFacade moze uzywac LINQ do filtrowania wolnych terminow."
    )

    doc.add_heading("12. Options Pattern", level=1)
    doc.add_paragraph(
        "Options Pattern porzadkuje konfiguracje Azure. Zamiast przekazywac surowe stringi po calej aplikacji, konfiguracja jest mapowana na klasy."
    )
    add_code_block(
        doc,
        "services.Configure<AzureBookingOptions>(configuration.GetSection(\"Azure\"));\n"
        "services.Configure<EmailOptions>(context.Configuration.GetSection(\"Email\"));",
        "Kod 11. Mapowanie konfiguracji na klasy Options.",
    )
    add_table(
        doc,
        ["Klasa", "Przechowywane ustawienia"],
        [
            ["AzureBookingOptions", "Cosmos, Storage, SignalR, Key Vault."],
            ["EmailOptions", "Connection string ACS Email i adres nadawcy."],
        ],
        [5.0, 11.0],
    )

    doc.add_heading("13. Funkcjonalnosc poparta wzorcami", level=1)
    add_table(
        doc,
        ["Funkcjonalnosc", "Uzyte wzorce", "Opis"],
        [
            ["Tworzenie wizyty", "Facade, Repository, Observer", "Facade koordynuje proces, Repository zapisuje dane, Event uruchamia reakcje."],
            ["Wysylka maila", "Observer, Handler, Adapter", "Po zdarzeniu handler wrzuca komunikat do kolejki, Function uzywa adaptera email."],
            ["Powiadomienia realtime", "Handler, Adapter", "SignalR handler wysyla komunikat do klientow online."],
            ["Zdjecia salonu/fryzjerow", "Adapter, Options Pattern", "IPhotoStorageService ukrywa szczegoly Blob Storage."],
            ["CRUD danych", "Repository, DI", "Kontrolery pracuja na IBookingRepository zamiast na CosmosClient."],
        ],
        [4.5, 4.5, 7.0],
    )
    add_screenshot_placeholder(
        doc,
        "Zrzut 1. Solution Explorer z projektami HairSalon.Booking.Api, Core i Functions.",
        "Wstaw zrzut ekranu z Visual Studio pokazujacy strukture projektu.",
    )
    add_screenshot_placeholder(
        doc,
        "Zrzut 2. Fragment kodu AppointmentBookingFacade w Visual Studio.",
        "Wstaw zrzut ekranu kodu wzorca Facade.",
    )
    add_screenshot_placeholder(
        doc,
        "Zrzut 3. Fragment kodu BookingEventPublisher lub IAppointmentBookedHandler.",
        "Wstaw zrzut ekranu kodu Observer/Handler.",
    )

    doc.add_heading("14. Przygotowanie projektu pod wzorce", level=1)
    doc.add_paragraph(
        "Projekt przygotowano tak, aby kazda klasa miala jedna odpowiedzialnosc. Kontrolery odpowiadaja za HTTP, Facade za proces biznesowy, Repository za dane, a handlery za reakcje po utworzeniu wizyty."
    )
    add_bullets(
        doc,
        [
            "Najpierw utworzono modele domenowe: Customer, Hairdresser, SalonService, Appointment, AppUser.",
            "Nastepnie utworzono interfejs IBookingRepository i implementacje CosmosBookingRepository.",
            "Proces rezerwacji przeniesiono do AppointmentBookingFacade.",
            "Reakcje po utworzeniu wizyty rozdzielono przez IAppointmentBookedHandler.",
            "Integracje Azure ukryto za adapterami i konfiguracja Options Pattern.",
            "Wszystkie zaleznosci zarejestrowano w ServiceCollectionExtensions.",
        ],
    )

    doc.add_heading("15. Podsumowanie", level=1)
    doc.add_paragraph(
        "Najwazniejsze wzorce w aplikacji to Facade, Repository, Observer/Event, Strategy/Handler i Factory Method. Pozostale wzorce i praktyki, takie jak Dependency Injection, Adapter, Singleton, Iterator i Options Pattern, wspieraja modularnosc aplikacji i integracje z Azure."
    )
    add_info_box(
        doc,
        "Gotowa odpowiedz na pytanie prowadzacego",
        "Wzorce zostaly zastosowane w praktycznych miejscach: Facade upraszcza rezerwacje, Repository oddziela Cosmos DB od logiki, Observer/Event uruchamia reakcje po utworzeniu wizyty, Strategy/Handler pozwala miec wiele reakcji na jedno zdarzenie, a Factory centralizuje tworzenie encji.",
        ORANGE,
    )

    path = ROOT / "Raport_2_Wzorce_Projektowe_HairSalonBooking.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    cloud = build_cloud_report()
    patterns = build_patterns_report()
    print(cloud)
    print(patterns)
