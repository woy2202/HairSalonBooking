from pathlib import Path
from zipfile import ZipFile
from docx import Document

for name in [
    "Raport_1_Chmura_Azure_HairSalonBooking.docx",
    "Raport_2_Wzorce_Projektowe_HairSalonBooking.docx",
]:
    path = Path(name)
    print(f"\n{name}")
    print(f"exists={path.exists()} size={path.stat().st_size if path.exists() else 0}")
    with ZipFile(path) as z:
        names = set(z.namelist())
        print(f"has_document_xml={'word/document.xml' in names}")
        print(f"media_files={len([n for n in names if n.startswith('word/media/')])}")
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    print(f"paragraphs={len(paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
    print(f"first_heading={paragraphs[0] if paragraphs else ''}")
