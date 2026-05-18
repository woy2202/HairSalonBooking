from __future__ import annotations

from pathlib import Path
from textwrap import dedent

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
IMAGE_DIR = ROOT / "report_code_screens"

BLUE = "1F4E79"
DARK = "111827"
GREEN = "D1FAE5"
GRAY = "F3F4F6"
LIGHT_BLUE = "EAF3FA"
YELLOW = "FFF7D6"
INK = "1F2937"


def set_cell_border(cell, color: str = "BFC7D1") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)

        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def write_cell(cell, text: str, bold: bool = False, color: str = INK, size: float = 9.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)


def style_doc(doc: Document, report_name: str) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in [("Title", 21), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.8)]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = report_name
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("6B7280")

    footer = section.footer.paragraphs[0]
    footer.text = "HairSalon Booking | raport projektowy"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor.from_string("6B7280")


def cover(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string("4B5563")

    doc.add_paragraph()
    heading = doc.add_paragraph()
    run = heading.add_run("Identyfikacja raportu")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    table = doc.add_table(rows=7, cols=2)
    rows = [
        ("Nazwa przedmiotu", "[wpisz nazwe przedmiotu]"),
        ("Grupa", "[wpisz grupe]"),
        ("Rok akademicki", "[wpisz rok akademicki]"),
        ("Student 1", "Wiktor Gochnio, nr indeksu: [wpisz numer indeksu]"),
        ("Student 2", "[imie i nazwisko], nr indeksu: [wpisz numer indeksu]"),
        ("Prowadzacy", "[wpisz prowadzacego, jesli wymagane]"),
        ("Data przygotowania", "18.05.2026"),
    ]

    for row, values in zip(table.rows, rows):
        write_cell(row.cells[0], values[0], True, BLUE)
        shade(row.cells[0], LIGHT_BLUE)
        write_cell(row.cells[1], values[1])

    doc.add_page_break()


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        write_cell(table.rows[0].cells[i], header, True, BLUE)
        shade(table.rows[0].cells[i], LIGHT_BLUE)

    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            write_cell(cells[i], value)

    doc.add_paragraph()


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def placeholder(doc: Document, caption: str, hint: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, GRAY)
    set_cell_border(cell, "9CA3AF")
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.paragraph_format.space_after = Pt(14)
    run = paragraph.add_run("[MIEJSCE NA ZRZUT EKRANU]\n")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    run = paragraph.add_run(hint + "\n\n\n")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("4B5563")

    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)
    doc.add_paragraph()


def get_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path(r"C:\Windows\Fonts\consola.ttf"),
        Path(r"C:\Windows\Fonts\cour.ttf"),
        Path(r"C:\Windows\Fonts\lucon.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)

    return ImageFont.load_default()


def snippet(path: str, start: int, end: int) -> str:
    file_path = ROOT / path
    lines = file_path.read_text(encoding="utf-8-sig").splitlines()
    selected = lines[start - 1:end]
    return "\n".join(f"{start + index:>3}  {line}" for index, line in enumerate(selected))


def code_image(name: str, file_label: str, code: str) -> Path:
    IMAGE_DIR.mkdir(exist_ok=True)
    font = get_font(18)
    small_font = get_font(15)
    lines = code.splitlines()
    line_height = 25
    padding = 24
    title_height = 34
    width = max(900, max(int(font.getlength(line)) for line in lines) + padding * 2)
    height = padding * 2 + title_height + max(1, len(lines)) * line_height

    image = Image.new("RGB", (width, height), "#111827")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, title_height), fill="#1F2937")
    draw.text((padding, 9), file_label, font=small_font, fill="#93C5FD")

    y = title_height + padding // 2
    for line in lines:
        number_part = line[:5]
        rest = line[5:]
        draw.text((padding, y), number_part, font=font, fill="#6B7280")
        draw.text((padding + 62, y), rest, font=font, fill="#D1FAE5")
        y += line_height

    out = IMAGE_DIR / f"{name}.png"
    image.save(out)
    return out


def add_code_screen(doc: Document, name: str, file_path: str, start: int, end: int, caption: str, explanation: str) -> None:
    code = snippet(file_path, start, end)
    img = code_image(name, file_path, code)
    doc.add_picture(str(img), width=Inches(6.7))
    last = doc.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)

    doc.add_paragraph(explanation)


def build_cloud_report() -> Path:
    doc = Document()
    style_doc(doc, "Raport chmurowy - Azure")
    cover(doc, "Raport 1: Chmura Azure", "HairSalon Booking - backend aplikacji salonu fryzjerskiego")

    doc.add_heading("1. Tytul projektu", level=1)
    doc.add_paragraph("HairSalon Booking - backend aplikacji salonu fryzjerskiego do umawiania wizyt.")

    doc.add_heading("2. Cel projektu", level=1)
    doc.add_paragraph(
        "Celem projektu bylo przygotowanie backendu aplikacji salonu fryzjerskiego dzialajacego w chmurze Microsoft Azure. "
        "System pokazuje, jak polaczyc aplikacje ASP.NET Core z uslugami App Service, Cosmos DB, Blob Storage, Queue Storage, Azure Functions, SignalR, Easy Auth, Key Vault, Monitor oraz Azure Communication Services Email."
    )

    doc.add_heading("3. Zalozenia projektu", level=1)
    doc.add_heading("3.1 Architektura", level=2)
    add_table(doc, ["Element", "Rola w projekcie", "Uzasadnienie"], [
        ["ASP.NET Core API", "Udostepnia kontrolery REST i Swagger", "Backend jest testowalny w Visual Studio oraz po publikacji w Azure App Service."],
        ["Azure App Service", "Hosting aplikacji booking-api", "Nie trzeba zarzadzac serwerem, a API jest dostepne przez publiczny adres HTTPS."],
        ["Azure Cosmos DB", "Baza dokumentowa", "Kontenery odpowiadaja glownym obiektom systemu: users, customers, barbers, services, appointments."],
        ["Azure Blob Storage", "Zdjecia i podsumowania wizyt", "Pliki binarne nie sa zapisywane w bazie, tylko w wyspecjalizowanym magazynie."],
        ["Azure Queue Storage", "Kolejka booking-notifications", "Oddziela zapis wizyty od wysylki maila."],
        ["Azure Functions", "Funkcja serverless", "Automatycznie odbiera komunikaty z kolejki i wysyla potwierdzenie."],
        ["Azure Communication Services Email", "Potwierdzenie mailowe", "Klient dostaje maila po rezerwacji wizyty."],
        ["Azure SignalR Service", "Powiadomienia realtime", "Otwarty klient testowy widzi natychmiast informacje o nowej wizycie."],
        ["Azure Key Vault", "Przechowywanie sekretow", "Connection stringi i klucze nie powinny byc wpisane w kodzie."],
        ["Azure Monitor / Application Insights", "Logi i diagnostyka", "Pozwala sprawdzac requesty, bledy, kolejke i dzialanie funkcji."],
        ["Easy Auth", "Logowanie Google/GitHub", "Azure zabezpiecza aplikacje przed wejscem do kodu backendu."],
    ])

    doc.add_heading("3.2 Przeplyw utworzenia wizyty", level=2)
    bullets(doc, [
        "Uzytkownik lub Swagger wysyla POST /api/Appointments.",
        "AppointmentsController buduje obiekt Appointment i przekazuje go do AppointmentBookingFacade.",
        "Fasada pobiera klienta, fryzjera i usluge z repozytoriow, sprawdza kolizje i zapisuje wizyte w Cosmos DB.",
        "Po zapisie publikowane jest zdarzenie AppointmentBooked.",
        "Handlery wysylaja podsumowanie do Blob Storage, komunikat do Queue Storage oraz powiadomienie SignalR.",
        "Azure Function odbiera komunikat z kolejki i wysyla mail przez Azure Communication Services Email.",
    ])

    doc.add_heading("4. Funkcjonalnosc", level=1)
    bullets(doc, [
        "CRUD klientow, fryzjerow, uslug i wizyt przez Swagger.",
        "Role uzytkownikow: klient, fryzjer, admin.",
        "Logowanie przez Easy Auth z providerami Google/GitHub.",
        "Automatyczne utworzenie klienta przy rejestracji uzytkownika.",
        "Admin moze zmieniac role uzytkownikow i przypisac konto fryzjera.",
        "Dodawanie zdjec salonu i fryzjerow do Blob Storage.",
        "Kolejka i Azure Function obsluguja maila w tle.",
        "SignalR pokazuje powiadomienia realtime po utworzeniu wizyty.",
    ])
    placeholder(doc, "Zrzut 1. Swagger UI z kontrolerami aplikacji.", "Wstaw ekran /swagger/index.html po zalogowaniu.")
    placeholder(doc, "Zrzut 2. Azure Portal - grupa zasobow Booking.", "Wstaw ekran z App Service, Cosmos DB, Function App, Storage, SignalR, ACS Email.")
    placeholder(doc, "Zrzut 3. Cosmos DB Data Explorer.", "Wstaw ekran z kontenerami users, customers, barbers, services, appointments.")
    placeholder(doc, "Zrzut 4. Mail z potwierdzeniem wizyty.", "Wstaw ekran otrzymanej wiadomosci od DoNotReply@...azurecomm.net.")

    doc.add_heading("5. Przygotowanie projektu", level=1)
    doc.add_heading("5.1 Kod startowy i konfiguracja chmury", level=2)
    add_code_screen(
        doc,
        "cloud_program",
        "HairSalon.Booking.Api/Program.cs",
        10,
        82,
        "Zrzut kodu 1. Program.cs w projekcie API.",
        "Program.cs ma teraz klasyczny zapis z metoda Main oraz metodami pomocniczymi. W tym miejscu aplikacja dolacza Key Vault, Application Insights, Swagger, CORS, kontrolery i SignalR Hub."
    )
    add_code_screen(
        doc,
        "cloud_di",
        "HairSalon.Booking.Api/Infrastructure/ServiceCollectionExtensions.cs",
        13,
        50,
        "Zrzut kodu 2. Rejestracja komponentow Azure i wzorcow w DI.",
        "Metoda AddBookingApplication pokazuje, jakie klasy sa wstrzykiwane do kontrolerow i serwisow. Tu aplikacja decyduje, czy uzyc Cosmos DB, czy lokalnego repozytorium in-memory."
    )

    doc.add_heading("5.2 Endpoint rezerwacji i logika biznesowa", level=2)
    add_code_screen(
        doc,
        "cloud_appointments_controller",
        "HairSalon.Booking.Api/Controllers/AppointmentsController.cs",
        8,
        76,
        "Zrzut kodu 3. AppointmentsController - endpointy CRUD wizyt.",
        "Kontroler ma klasyczny konstruktor i prywatne pola. POST tworzy obiekt Appointment i przekazuje go do fasady, zamiast trzymac cala logike w kontrolerze."
    )
    add_code_screen(
        doc,
        "cloud_booking_facade",
        "HairSalon.Booking.Core/Services/AppointmentBookingFacade.cs",
        9,
        61,
        "Zrzut kodu 4. AppointmentBookingFacade - glowny proces rezerwacji.",
        "Fasada pobiera dane z repozytoriow, sprawdza warunki, liczy czas zakonczenia, zapisuje wizyte i publikuje zdarzenie dla komponentow Azure."
    )

    doc.add_heading("5.3 Cosmos DB i Blob Storage", level=2)
    add_code_screen(
        doc,
        "cloud_cosmos_repository",
        "HairSalon.Booking.Api/Data/CosmosBookingRepository.cs",
        7,
        55,
        "Zrzut kodu 5. CosmosBookingRepository - zapis i odczyt dokumentow.",
        "Repozytorium ukrywa SDK Cosmos DB. Dla partition key /id metody ReadItemAsync, CreateItemAsync i UpsertItemAsync uzywaja new PartitionKey(id)."
    )
    add_code_screen(
        doc,
        "cloud_blob_photo",
        "HairSalon.Booking.Api/Infrastructure/BlobPhotoStorageService.cs",
        8,
        60,
        "Zrzut kodu 6. BlobPhotoStorageService - upload zdjec.",
        "Serwis waliduje typ pliku, tworzy kontener Blob Storage, wysyla plik i zwraca nazwe oraz URL bloba."
    )

    doc.add_heading("5.4 Queue, Function, Email i SignalR", level=2)
    add_code_screen(
        doc,
        "cloud_queue_handler",
        "HairSalon.Booking.Api/Infrastructure/AzureQueueAppointmentBookedHandler.cs",
        12,
        80,
        "Zrzut kodu 7. AzureQueueAppointmentBookedHandler - komunikat do kolejki.",
        "Handler po utworzeniu wizyty buduje JSON z danymi wizyty i wysyla go do kolejki booking-notifications."
    )
    add_code_screen(
        doc,
        "cloud_function",
        "HairSalon.Booking.Functions/AppointmentBookedQueueFunction.cs",
        8,
        74,
        "Zrzut kodu 8. AppointmentBookedQueueFunction - reakcja serverless.",
        "Azure Function uruchamia sie po pojawieniu wiadomosci w kolejce. Deserializuje komunikat i przekazuje go do serwisu mailowego."
    )
    add_code_screen(
        doc,
        "cloud_email",
        "HairSalon.Booking.Functions/Email/AzureCommunicationEmailSender.cs",
        8,
        58,
        "Zrzut kodu 9. AzureCommunicationEmailSender - ACS Email.",
        "Serwis tworzy EmailClient na podstawie konfiguracji i wysyla do klienta potwierdzenie rezerwacji."
    )
    add_code_screen(
        doc,
        "cloud_signalr",
        "HairSalon.Booking.Api/Infrastructure/SignalRAppointmentBookedHandler.cs",
        8,
        38,
        "Zrzut kodu 10. SignalRAppointmentBookedHandler - realtime.",
        "SignalR wysyla wiadomosc appointmentBooked do wszystkich klientow oraz hairdresserAppointmentBooked do grupy danego fryzjera."
    )

    doc.add_heading("6. Podsumowanie", level=1)
    doc.add_paragraph(
        "Projekt pokazuje kompletna aplikacje backendowa w Azure. Najwazniejszy przeplyw techniczny laczy App Service, Cosmos DB, Queue Storage, Azure Functions i ACS Email. "
        "Dodatkowo zastosowano Blob Storage, SignalR, Easy Auth, Key Vault oraz Application Insights, co dobrze pokazuje praktyczne uzycie chmury w projekcie uczelnianym."
    )

    path = ROOT / "Raport_Chmura_Azure_HairSalonBooking_aktualny.docx"
    doc.save(path)
    return path


def build_patterns_report() -> Path:
    doc = Document()
    style_doc(doc, "Raport wzorcow projektowych")
    cover(doc, "Raport 2: Wzorce projektowe", "HairSalon Booking - aktualny kod po refaktoryzacji stylistycznej")

    doc.add_heading("1. Tytul projektu", level=1)
    doc.add_paragraph("HairSalon Booking - backend aplikacji salonu fryzjerskiego do umawiania wizyt.")

    doc.add_heading("2. Cel projektu", level=1)
    doc.add_paragraph(
        "Celem raportu jest pokazanie, jakie wzorce projektowe zostaly uzyte w aktualnym kodzie aplikacji oraz jak widac je w Visual Studio. "
        "Kod zostal przepisany w bardziej klasycznym stylu C#, zgodnym z przykladami z prezentacji: jawne konstruktory, klasy zamiast rekordow DTO, blokowe namespace oraz klasyczne zdarzenia."
    )

    doc.add_heading("3. Zalozenia projektu", level=1)
    add_table(doc, ["Warstwa", "Odpowiedzialnosc", "Uzyte wzorce"], [
        ["HairSalon.Booking.Api", "Kontrolery, konfiguracja, integracje Azure", "Dependency Injection, Adapter, Repository"],
        ["HairSalon.Booking.Core", "Modele, logika biznesowa, zdarzenia", "Facade, Factory Method, Observer/Event, Iterator"],
        ["HairSalon.Booking.Functions", "Przetwarzanie kolejki i email", "Adapter, Options Pattern, event-driven processing"],
    ])
    bullets(doc, [
        "Kontrolery nie tworza bezposrednio klientow Azure SDK.",
        "Dostep do bazy ukrywa interfejs IBookingRepository.",
        "Proces rezerwacji jest zebrany w AppointmentBookingFacade.",
        "Po utworzeniu wizyty aplikacja publikuje zdarzenie, a osobne handlery reaguja niezaleznie.",
        "Fabryka tworzy typy domenowe w jednym miejscu.",
        "Iterator udostepnia sloty dostepnosci jako kolekcje IEnumerable.",
    ])

    doc.add_heading("4. Funkcjonalnosc", level=1)
    add_table(doc, ["Funkcjonalnosc", "Wzorzec", "Jak pomaga"], [
        ["Tworzenie wizyty", "Facade", "Kontroler wywoluje jedna metode BookAsync zamiast znac caly proces rezerwacji."],
        ["CRUD w Cosmos DB", "Repository", "Kontrolery korzystaja z interfejsu, a nie z Cosmos SDK."],
        ["Reakcje po rezerwacji", "Observer/Event + Handler", "Mail, SignalR i Blob Storage sa uruchamiane jako oddzielne reakcje."],
        ["Tworzenie encji", "Factory Method", "Jeden punkt tworzenia obiektow domenowych."],
        ["Wolne terminy", "Iterator", "Sloty terminow sa przechodzone jak zwykla kolekcja."],
        ["Azure Blob/Queue/Email/SignalR", "Adapter", "Szczegoly SDK Azure sa ukryte za klasami infrastruktury."],
        ["Konfiguracja sekretow", "Options Pattern", "Klasy pobieraja opcje z konfiguracji zamiast czytac recznie appsettings."],
    ])
    placeholder(doc, "Zrzut 1. Solution Explorer z projektami Api, Core i Functions.", "Wstaw ekran Visual Studio z aktualna struktura projektu.")
    placeholder(doc, "Zrzut 2. Swagger pokazujacy kontrolery CRUD.", "Wstaw ekran z endpointami Customers, Hairdressers, Appointments, SalonServices.")

    doc.add_heading("5. Przygotowanie projektu", level=1)
    doc.add_heading("5.1 Facade", level=2)
    add_code_screen(
        doc,
        "pattern_facade",
        "HairSalon.Booking.Core/Services/AppointmentBookingFacade.cs",
        9,
        61,
        "Zrzut kodu 1. Facade - AppointmentBookingFacade.",
        "Facade udostepnia jedna metode BookAsync. W srodku koordynuje pobranie danych, walidacje, sprawdzenie terminu, zapis i publikacje zdarzenia."
    )

    doc.add_heading("5.2 Repository", level=2)
    add_code_screen(
        doc,
        "pattern_repository_interface",
        "HairSalon.Booking.Core/Repositories/IBookingRepository.cs",
        1,
        18,
        "Zrzut kodu 2. Interfejs Repository.",
        "Interfejs okresla kontrakt dla operacji CRUD. Kontrolery moga korzystac z repozytorium bez wiedzy, czy dane sa w Cosmos DB, czy w pamieci."
    )
    add_code_screen(
        doc,
        "pattern_repository_cosmos",
        "HairSalon.Booking.Api/Data/CosmosBookingRepository.cs",
        7,
        55,
        "Zrzut kodu 3. Implementacja Repository dla Cosmos DB.",
        "CosmosBookingRepository implementuje IBookingRepository i ukrywa szczegoly Azure Cosmos DB."
    )

    doc.add_heading("5.3 Observer/Event oraz Handler", level=2)
    add_code_screen(
        doc,
        "pattern_event",
        "HairSalon.Booking.Core/Events/BookingEventPublisher.cs",
        4,
        33,
        "Zrzut kodu 4. Observer/Event - BookingEventPublisher.",
        "Publisher posiada event AppointmentBooked i liste handlerow. Po utworzeniu wizyty powiadamia zainteresowane elementy systemu."
    )
    add_code_screen(
        doc,
        "pattern_handler",
        "HairSalon.Booking.Core/Events/IAppointmentBookedHandler.cs",
        1,
        10,
        "Zrzut kodu 5. Handler zdarzenia wizyty.",
        "Interfejs pozwala dodawac kolejne reakcje na utworzenie wizyty bez przepisywania fasady."
    )
    add_table(doc, ["Handler", "Efekt po utworzeniu wizyty"], [
        ["AzureBlobAppointmentSummaryHandler", "Tworzy plik tekstowy z podsumowaniem wizyty w Blob Storage."],
        ["AzureQueueAppointmentBookedHandler", "Wysyla komunikat do booking-notifications."],
        ["SignalRAppointmentBookedHandler", "Wysyla powiadomienie realtime do klientow SignalR."],
    ])

    doc.add_heading("5.4 Factory Method", level=2)
    add_code_screen(
        doc,
        "pattern_factory",
        "HairSalon.Booking.Core/Patterns/BookingEntityFactory.cs",
        4,
        34,
        "Zrzut kodu 6. Factory Method - BookingEntityFactory.",
        "Fabryka tworzy obiekty domenowe na podstawie EntityKind. Zastosowano klasyczny switch, podobny do przykladow z prezentacji."
    )

    doc.add_heading("5.5 Iterator", level=2)
    add_code_screen(
        doc,
        "pattern_iterator",
        "HairSalon.Booking.Core/Patterns/AvailabilitySlotCollection.cs",
        5,
        30,
        "Zrzut kodu 7. Iterator - AvailabilitySlotCollection.",
        "Kolekcja slotow implementuje IEnumerable<DateTimeOffset>, wiec mozna po niej przechodzic w petlach i zapytaniach LINQ bez ujawniania szczegolow listy."
    )

    doc.add_heading("5.6 Dependency Injection, Adapter i Options Pattern", level=2)
    add_code_screen(
        doc,
        "pattern_di",
        "HairSalon.Booking.Api/Infrastructure/ServiceCollectionExtensions.cs",
        15,
        50,
        "Zrzut kodu 8. Dependency Injection.",
        "Kontener DI laczy interfejsy z implementacjami. Dzieki temu kod jest bardziej uporzadkowany i latwiej wymienic implementacje."
    )
    add_code_screen(
        doc,
        "pattern_adapter_blob",
        "HairSalon.Booking.Api/Infrastructure/BlobPhotoStorageService.cs",
        8,
        60,
        "Zrzut kodu 9. Adapter dla Blob Storage.",
        "BlobPhotoStorageService ukrywa Azure SDK za IPhotoStorageService, czyli kontroler nie musi znac szczegolow BlobClient."
    )
    add_code_screen(
        doc,
        "pattern_options",
        "HairSalon.Booking.Api/Options/AzureBookingOptions.cs",
        1,
        41,
        "Zrzut kodu 10. Options Pattern - konfiguracja Azure.",
        "Opcje przechowuja nazwy kontenerow, connection stringi i nazwy kolejek. Klasy infrastruktury pobieraja je przez IOptions."
    )

    doc.add_heading("6. Podsumowanie", level=1)
    doc.add_paragraph(
        "Aktualny kod aplikacji zawiera kilka wzorcow omawianych na prezentacji: Factory Method, Facade, Iterator oraz Event/Observer. "
        "Dodatkowo zastosowano typowe wzorce aplikacyjne .NET: Repository, Dependency Injection, Adapter i Options Pattern. "
        "Po refaktoryzacji kod jest zapisany bardziej klasycznie, co ulatwia wskazanie elementow wzorcow podczas omawiania projektu."
    )

    path = ROOT / "Raport_Wzorce_Projektowe_HairSalonBooking_aktualny.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_cloud_report())
    print(build_patterns_report())
