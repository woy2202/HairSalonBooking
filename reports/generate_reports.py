from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import textwrap

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports"
ASSETS = OUT / "assets"
ASSETS.mkdir(parents=True, exist_ok=True)

ACCENT = RGBColor(38, 90, 142)
MUTED = RGBColor(90, 97, 110)


def load_font(path, size):
    try:
        return ImageFont.truetype(path, size)
    except OSError:
        return ImageFont.load_default()


MONO = load_font(r"C:\Windows\Fonts\consola.ttf", 22)
MONO_BOLD = load_font(r"C:\Windows\Fonts\consolab.ttf", 22)
TEXT = load_font(r"C:\Windows\Fonts\segoeui.ttf", 27)
TEXT_BOLD = load_font(r"C:\Windows\Fonts\segoeuib.ttf", 34)
SMALL = load_font(r"C:\Windows\Fonts\segoeui.ttf", 22)


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_doc(doc):
    doc.styles["Normal"].font.name = "Segoe UI"
    doc.styles["Normal"].font.size = Pt(10)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.name = "Segoe UI"
    doc.styles["Heading 1"].font.color.rgb = ACCENT
    doc.styles["Heading 1"].font.size = Pt(17)
    doc.styles["Heading 2"].font.color.rgb = RGBColor(44, 62, 80)
    doc.styles["Heading 2"].font.size = Pt(13)
    for section in doc.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)


def title(doc, main, sub):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(main)
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = ACCENT
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(sub)
    r.font.size = Pt(11)
    r.font.color.rgb = MUTED
    doc.add_paragraph()


def identification(doc):
    doc.add_heading("Identyfikacja raportu", level=1)
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ("Nazwa przedmiotu", "........................................................"),
        ("Grupa", "........................................................"),
        ("Rok akademicki", "........................................................"),
        ("Imiona i nazwiska studentów", "........................................................"),
        ("Numery indeksów", "........................................................"),
    ]
    for row, data in zip(table.rows, rows):
        cell_text(row.cells[0], data[0], bold=True)
        cell_text(row.cells[1], data[1])
        shade(row.cells[0], "EAF1F8")
    doc.add_paragraph()


def note(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, "F3F7FB")
    run = cell.paragraphs[0].add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(48, 66, 84)
    doc.add_paragraph()


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED


def add_image(doc, image_path, cap, width=6.6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(width))
    caption(doc, cap)


def placeholder(name, heading, body):
    path = ASSETS / f"{name}.png"
    img = Image.new("RGB", (1500, 620), "#F6F8FB")
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle((20, 20, 1480, 600), radius=18, outline="#8AA4C0", width=4, fill="#FFFFFF")
    draw.text((60, 60), heading, fill="#265A8E", font=TEXT_BOLD)
    y = 140
    for line in textwrap.wrap(body, width=92):
        draw.text((60, y), line, fill="#333A45", font=TEXT)
        y += 42
    draw.rectangle((60, 510, 1440, 560), outline="#C9D6E3", width=2)
    draw.text((82, 522), "Miejsce na wklejenie własnego zrzutu ekranu", fill="#6B7280", font=SMALL)
    img.save(path)
    return path


def diagram(name, heading, boxes, arrows):
    path = ASSETS / f"{name}.png"
    img = Image.new("RGB", (1600, 900), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw.text((50, 30), heading, fill="#1F4E79", font=TEXT_BOLD)
    pos = {}
    for key, label, x, y, w, h, color in boxes:
        pos[key] = (x, y, w, h)
        draw.rounded_rectangle((x, y, x + w, y + h), radius=20, fill=color, outline="#23415C", width=3)
        yy = y + 24
        for line in textwrap.wrap(label, width=max(14, w // 18)):
            tw = draw.textlength(line, font=TEXT)
            draw.text((x + (w - tw) / 2, yy), line, fill="#102033", font=TEXT)
            yy += 38
    for src, dst, label in arrows:
        x1, y1, w1, h1 = pos[src]
        x2, y2, w2, h2 = pos[dst]
        start = (x1 + w1, y1 + h1 / 2)
        end = (x2, y2 + h2 / 2)
        if x2 < x1:
            start = (x1, y1 + h1 / 2)
            end = (x2 + w2, y2 + h2 / 2)
        draw.line((start, end), fill="#52677A", width=4)
        ax, ay = end
        if end[0] > start[0]:
            draw.polygon([(ax, ay), (ax - 16, ay - 9), (ax - 16, ay + 9)], fill="#52677A")
        else:
            draw.polygon([(ax, ay), (ax + 16, ay - 9), (ax + 16, ay + 9)], fill="#52677A")
        if label:
            mx, my = (start[0] + end[0]) / 2, (start[1] + end[1]) / 2
            draw.rounded_rectangle((mx - 145, my - 24, mx + 145, my + 24), radius=10, fill="#F7FAFD", outline="#C8D4DF")
            tw = draw.textlength(label, font=SMALL)
            draw.text((mx - tw / 2, my - 15), label, fill="#31485C", font=SMALL)
    img.save(path)
    return path


def code_image(name, title_text, rel, start, end):
    path = ASSETS / f"{name}.png"
    lines = (ROOT / rel).read_text(encoding="utf-8", errors="replace").splitlines()
    selected = [f"{i:>3}  {lines[i - 1]}" for i in range(start, min(end, len(lines)) + 1)]
    max_len = max([len(x) for x in selected] + [len(title_text)])
    width = min(1900, max(1000, max_len * 13 + 90))
    height = 95 + len(selected) * 34 + 35
    img = Image.new("RGB", (width, height), "#111827")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, 62), fill="#1F2937")
    draw.text((28, 17), title_text, fill="#E5E7EB", font=MONO_BOLD)
    y = 82
    for line in selected:
        draw.text((28, y), line[:150], fill="#D1D5DB", font=MONO)
        y += 34
    img.save(path)
    return path


def component_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(table.rows[0].cells, ["Komponent/wzorzec", "Rola w projekcie", "Uzasadnienie"]):
        cell_text(cell, header, bold=True, color=RGBColor(255, 255, 255))
        shade(cell, "265A8E")
    for data in rows:
        cells = table.add_row().cells
        for i, text in enumerate(data):
            cell_text(cells[i], text)
        shade(cells[0], "EAF1F8")
    doc.add_paragraph()


def build_assets():
    assets = {}
    assets["azure_arch"] = diagram("azure_architecture", "Architektura chmurowa backendu HairSalon Booking", [
        ("client", "Klient / Admin / Fryzjer\nSwagger lub aplikacja", 50, 170, 300, 110, "#D9EAF7"),
        ("api", "Azure App Service\nHairSalon.Booking.Api", 450, 150, 330, 150, "#DDF4E8"),
        ("auth", "Easy Auth\nGoogle / GitHub", 450, 370, 330, 110, "#FFF3D7"),
        ("kv", "Azure Key Vault\nsekrety", 880, 90, 330, 120, "#FDE2E2"),
        ("cosmos", "Azure Cosmos DB\nkontenery danych", 880, 260, 330, 130, "#E8E3FF"),
        ("storage", "Storage Account\nBlob + Queue", 880, 450, 330, 130, "#E4F2FF"),
        ("func", "Azure Function\nobsługa kolejki", 1250, 450, 300, 130, "#FCE7F3"),
        ("email", "ACS Email\npotwierdzenia", 1250, 630, 300, 110, "#E7F8EA"),
        ("signalr", "Azure SignalR\nrealtime", 1250, 250, 300, 120, "#F1F5F9"),
        ("monitor", "Application Insights\nlogi", 880, 650, 330, 110, "#FFF7ED"),
    ], [
        ("client", "api", "HTTP"), ("api", "auth", "nagłówki"), ("api", "kv", "sekrety"),
        ("api", "cosmos", "CRUD"), ("api", "storage", "blob/queue"), ("storage", "func", "trigger"),
        ("func", "email", "send"), ("api", "signalr", "push"), ("api", "monitor", "telemetria"),
    ])
    assets["patterns_arch"] = diagram("patterns_architecture", "Przepływ wzorców w kodzie aplikacji", [
        ("controller", "Controllers\nAPI endpoints", 50, 190, 300, 120, "#D9EAF7"),
        ("facade", "Facade\nAppointmentBookingFacade", 440, 170, 340, 150, "#DDF4E8"),
        ("repo", "Repository\nIBookingRepository<T>", 880, 140, 330, 130, "#E8E3FF"),
        ("cosmos", "CosmosRepository\nimplementacja", 1240, 140, 300, 130, "#F1F5F9"),
        ("factory", "Factory\nBookingEntityFactory", 880, 330, 330, 120, "#FFF3D7"),
        ("publisher", "Observer / Publisher\nBookingEventPublisher", 440, 430, 340, 140, "#FDE2E2"),
        ("queue", "Handler\nAzure Queue", 880, 540, 330, 120, "#E4F2FF"),
        ("signalr", "Handler\nSignalR", 1240, 540, 300, 120, "#FCE7F3"),
        ("di", "Dependency Injection\nServiceCollection", 50, 520, 300, 120, "#E7F8EA"),
    ], [
        ("controller", "facade", "book"), ("facade", "repo", "dane"), ("repo", "cosmos", "adapter"),
        ("facade", "publisher", "event"), ("publisher", "queue", "notify"), ("publisher", "signalr", "notify"),
        ("di", "controller", "wstrzykuje"), ("di", "repo", "rejestruje"),
    ])
    placeholders = {
        "rg": ("Screen: Resource Group Booking", "Wklej widok grupy zasobów Azure z App Service, Cosmos DB, Function App, Storage Account, SignalR, Key Vault, Application Insights i Email Communication Service."),
        "cosmos": ("Screen: Cosmos DB Data Explorer", "Wklej kontenery: users, customers, barbers, services, appointments, reviews i salon."),
        "swagger": ("Screen: Swagger API", "Wklej listę kontrolerów i przykładowy wynik CRUD."),
        "logs": ("Screen: Function App Log Stream", "Wklej logi po utworzeniu wizyty i wysłaniu wiadomości e-mail."),
        "signalr": ("Screen: test SignalR", "Wklej stronę testową pokazującą komunikat appointmentBooked."),
        "vs": ("Screen: Visual Studio Solution Explorer", "Wklej widok trzech projektów: Api, Core i Functions."),
        "flow": ("Screen: test rezerwacji w Swagger", "Wklej wynik POST /api/Appointments i odczyt utworzonej wizyty."),
    }
    for key, value in placeholders.items():
        assets[f"ph_{key}"] = placeholder(f"ph_{key}", value[0], value[1])
    code_specs = {
        "program_kv": ("Program.cs - Key Vault i konfiguracja startowa", "HairSalon.Booking.Api/Program.cs", 13, 33),
        "program_pipeline": ("Program.cs - middleware i endpoint SignalR", "HairSalon.Booking.Api/Program.cs", 78, 91),
        "services_di": ("ServiceCollectionExtensions.cs - DI, SignalR, Cosmos", "HairSalon.Booking.Api/Infrastructure/ServiceCollectionExtensions.cs", 15, 52),
        "cosmos_repo": ("CosmosBookingRepository.cs - CRUD w Cosmos DB", "HairSalon.Booking.Api/Data/CosmosBookingRepository.cs", 15, 66),
        "cosmos_resolver": ("CosmosContainerResolver.cs - encje i kontenery", "HairSalon.Booking.Api/Data/CosmosContainerResolver.cs", 18, 75),
        "blob_photo": ("BlobPhotoStorageService.cs - upload zdjęć", "HairSalon.Booking.Api/Infrastructure/BlobPhotoStorageService.cs", 24, 80),
        "queue_handler": ("AzureQueueAppointmentBookedHandler.cs - Storage Queue", "HairSalon.Booking.Api/Infrastructure/AzureQueueAppointmentBookedHandler.cs", 35, 73),
        "function_queue": ("AppointmentBookedQueueFunction.cs - QueueTrigger", "HairSalon.Booking.Functions/AppointmentBookedQueueFunction.cs", 26, 50),
        "email_sender": ("AzureCommunicationEmailSender.cs - ACS Email", "HairSalon.Booking.Functions/Email/AzureCommunicationEmailSender.cs", 23, 47),
        "signalr_handler": ("SignalRAppointmentBookedHandler.cs - realtime", "HairSalon.Booking.Api/Infrastructure/SignalRAppointmentBookedHandler.cs", 16, 35),
        "easyauth": ("EasyAuthCurrentUserService.cs - Google/GitHub", "HairSalon.Booking.Api/Infrastructure/EasyAuthCurrentUserService.cs", 22, 51),
        "repo_interface": ("IBookingRepository.cs - kontrakt", "HairSalon.Booking.Core/Repositories/IBookingRepository.cs", 3, 12),
        "factory": ("BookingEntityFactory.cs - Factory", "HairSalon.Booking.Core/Patterns/BookingEntityFactory.cs", 5, 36),
        "facade": ("AppointmentBookingFacade.cs - Facade", "HairSalon.Booking.Core/Services/AppointmentBookingFacade.cs", 30, 57),
        "facade_validation": ("AppointmentBookingFacade.cs - walidacje", "HairSalon.Booking.Core/Services/AppointmentBookingFacade.cs", 75, 117),
        "publisher": ("BookingEventPublisher.cs - Observer/Publisher", "HairSalon.Booking.Core/Events/BookingEventPublisher.cs", 5, 31),
        "iterator": ("AvailabilitySlotCollection.cs - Iterator", "HairSalon.Booking.Core/Patterns/AvailabilitySlotCollection.cs", 5, 28),
        "auth_register": ("AuthController.cs - konto user + customer", "HairSalon.Booking.Api/Controllers/AuthController.cs", 55, 96),
    }
    for key, spec in code_specs.items():
        assets[key] = code_image(f"code_{key}", *spec)
    return assets


def cloud_report(assets):
    doc = Document()
    style_doc(doc)
    title(doc, "Raport prezentujący projekt - Azure", "HairSalon Booking | Wymagana akceptacja, nie podlega ocenie")
    identification(doc)
    doc.add_heading("Tytuł projektu", level=1)
    doc.add_paragraph("Backend aplikacji do umawiania wizyt w salonie fryzjerskim z wykorzystaniem usług Microsoft Azure.")
    doc.add_heading("Cel projektu", level=1)
    doc.add_paragraph("Celem projektu było stworzenie backendu w C#/.NET dla salonu fryzjerskiego. Aplikacja umożliwia zarządzanie klientami, fryzjerami, usługami, wizytami, opiniami oraz zdjęciami. Część chmurowa pokazuje praktyczne użycie Azure w realnym przepływie biznesowym.")
    note(doc, "Najważniejszy przepływ: klient tworzy wizytę w API, dane trafiają do Cosmos DB, zdjęcia do Blob Storage, zdarzenie do kolejki, Azure Function wysyła e-mail, a SignalR informuje użytkowników realtime.")
    doc.add_heading("Założenia projektu", level=1)
    doc.add_paragraph("Rozwiązanie składa się z trzech projektów: HairSalon.Booking.Api, HairSalon.Booking.Core oraz HairSalon.Booking.Functions. API jest hostowane w App Service, Core zawiera logikę domenową, a Functions obsługuje zadania asynchroniczne.")
    add_image(doc, assets["azure_arch"], "Rysunek 1. Architektura chmurowa aplikacji.", 6.7)
    add_image(doc, assets["ph_vs"], "Rysunek 2. Miejsce na screen rozwiązania w Visual Studio.", 6.7)
    doc.add_heading("Wybrane technologie - uzasadnienie", level=2)
    component_table(doc, [
        ("App Service", "Hostuje API i Swagger.", "Prosty deployment ASP.NET Core."),
        ("Cosmos DB", "Przechowuje dokumenty aplikacji.", "Elastyczny JSON i kontenery dla encji."),
        ("Blob Storage", "Przechowuje zdjęcia salonu i fryzjerów.", "Pliki nie obciążają bazy danych."),
        ("Queue Storage", "Buforuje zdarzenia rezerwacji.", "Oddziela API od wysyłki maila."),
        ("Azure Functions", "Odbiera wiadomości z kolejki.", "Serverless dla krótkich zadań."),
        ("ACS Email", "Wysyła potwierdzenia wizyt.", "Realna usługa komunikacji."),
        ("SignalR", "Powiadomienia realtime.", "Admin i fryzjer widzą nowe wizyty bez odświeżania."),
        ("Key Vault", "Sekrety poza kodem.", "Bezpieczniejsza konfiguracja."),
        ("Application Insights", "Logi i telemetry.", "Diagnoza requestów i działania funkcji."),
        ("Easy Auth", "Google/GitHub login.", "Azure blokuje niezalogowanych przed API."),
    ])
    doc.add_heading("Kod uruchamiający komponenty Azure", level=1)
    add_image(doc, assets["program_kv"], "Rysunek 3. Key Vault jako źródło konfiguracji.", 6.7)
    add_image(doc, assets["program_pipeline"], "Rysunek 4. Pipeline API i endpoint huba SignalR.", 6.7)
    add_image(doc, assets["services_di"], "Rysunek 5. Rejestracja Cosmos, SignalR, handlerów i serwisów.", 6.7)
    doc.add_heading("Funkcjonalność", level=1)
    bullets(doc, [
        "Logowanie przez Google/GitHub i utworzenie konta w users oraz customers.",
        "CRUD klientów, fryzjerów, usług, wizyt, opinii oraz zdjęć.",
        "Panel admina: zarządzanie użytkownikami, rolami, usługami i danymi.",
        "Klient może edytować swoje dane, odwołać wizytę i usunąć konto.",
        "Fryzjer widzi swoje wizyty, historię klienta i zmienia status wizyty.",
        "Admin nadaje klientowi rolę fryzjera, a aplikacja przenosi profil w bazie.",
        "Po rezerwacji wysyłany jest e-mail i powiadomienie realtime.",
    ])
    add_image(doc, assets["ph_swagger"], "Rysunek 6. Miejsce na screen Swaggera z endpointami.", 6.7)
    add_image(doc, assets["ph_rg"], "Rysunek 7. Miejsce na screen Resource Group w Azure.", 6.7)
    doc.add_heading("Cosmos DB", level=1)
    doc.add_paragraph("Cosmos DB jest główną bazą danych. Każdy model domenowy ma swój kontener, a generyczne repozytorium udostępnia CRUD. Kontrolery nie używają bezpośrednio Cosmos SDK.")
    add_image(doc, assets["cosmos_resolver"], "Rysunek 8. Mapowanie encji na kontenery Cosmos DB.", 6.7)
    add_image(doc, assets["cosmos_repo"], "Rysunek 9. Generyczny CRUD w Cosmos DB.", 6.7)
    add_image(doc, assets["ph_cosmos"], "Rysunek 10. Miejsce na screen Data Explorer.", 6.7)
    doc.add_heading("Blob Storage", level=1)
    doc.add_paragraph("Blob Storage przechowuje zdjęcia. API waliduje typ pliku, rozmiar, zapisuje blob i zwraca adres. W Cosmos DB pozostają metadane zdjęcia.")
    add_image(doc, assets["blob_photo"], "Rysunek 11. Upload zdjęcia do Azure Blob Storage.", 6.7)
    doc.add_heading("Queue Storage, Azure Functions i Email", level=1)
    doc.add_paragraph("Po utworzeniu wizyty API zapisuje komunikat w kolejce. Azure Function odbiera wiadomość i wysyła potwierdzenie przez Azure Communication Services Email.")
    add_image(doc, assets["queue_handler"], "Rysunek 12. Komunikat rezerwacji wysyłany do kolejki.", 6.7)
    add_image(doc, assets["function_queue"], "Rysunek 13. QueueTrigger uruchamiający funkcję serverless.", 6.7)
    add_image(doc, assets["email_sender"], "Rysunek 14. Wysyłanie potwierdzenia przez ACS Email.", 6.7)
    add_image(doc, assets["ph_logs"], "Rysunek 15. Miejsce na screen logów Function App.", 6.7)
    doc.add_heading("SignalR i Easy Auth", level=1)
    add_image(doc, assets["signalr_handler"], "Rysunek 16. SignalR wysyła komunikat do admina i fryzjera.", 6.7)
    add_image(doc, assets["ph_signalr"], "Rysunek 17. Miejsce na screen testu SignalR.", 6.7)
    add_image(doc, assets["easyauth"], "Rysunek 18. Dane użytkownika z nagłówków Easy Auth.", 6.7)
    add_image(doc, assets["auth_register"], "Rysunek 19. Rejestracja tworzy users i customers.", 6.7)
    doc.add_heading("Przygotowanie projektu", level=1)
    doc.add_paragraph("Projekt powstawał etapami: najpierw modele i kontrolery, potem repozytorium Cosmos DB, następnie Blob Storage, kolejki, Functions, e-mail, SignalR, Key Vault i monitoring.")
    numbered(doc, [
        "Utworzenie rozwiązania i projektów Api/Core/Functions.",
        "Dodanie modeli domenowych i kontrolerów CRUD.",
        "Dodanie Cosmos DB oraz kontenerów.",
        "Dodanie Blob Storage dla zdjęć.",
        "Dodanie kolejki booking-notifications i Azure Function.",
        "Dodanie ACS Email do potwierdzeń.",
        "Dodanie SignalR do realtime.",
        "Dodanie Key Vault i Application Insights.",
        "Publikacja do Azure i test w Swaggerze.",
    ])
    add_image(doc, assets["ph_flow"], "Rysunek 20. Miejsce na screen testu utworzenia wizyty.", 6.7)
    path = OUT / "Raport_Azure_HairSalonBooking.docx"
    doc.save(path)
    return path


def patterns_report(assets):
    doc = Document()
    style_doc(doc)
    title(doc, "Raport prezentujący projekt - wzorce projektowe", "HairSalon Booking | Wymagana akceptacja, nie podlega ocenie")
    identification(doc)
    doc.add_heading("Tytuł projektu", level=1)
    doc.add_paragraph("Backend aplikacji salonu fryzjerskiego z wykorzystaniem wzorców projektowych w C#/.NET.")
    doc.add_heading("Cel projektu", level=1)
    doc.add_paragraph("Celem raportu jest pokazanie, że backend ma uporządkowaną strukturę, a wzorce wynikają z realnych potrzeb: zapisu danych, rezerwacji wizyt, tworzenia encji, zdarzeń i integracji z usługami zewnętrznymi.")
    doc.add_heading("Założenia projektu", level=1)
    doc.add_paragraph("Wzorce zastosowano w warstwie Core i Api. Kontrolery pozostają możliwie proste, a logika biznesowa i infrastrukturalna jest przeniesiona do serwisów, repozytoriów i handlerów.")
    add_image(doc, assets["patterns_arch"], "Rysunek 1. Przepływ wzorców w aplikacji.", 6.7)
    component_table(doc, [
        ("Repository", "IBookingRepository<T> i CosmosBookingRepository<T>.", "Kontrolery nie zależą od konkretnej bazy."),
        ("Factory", "BookingEntityFactory.", "Tworzy właściwy typ encji domenowej."),
        ("Facade", "AppointmentBookingFacade.", "Ukrywa złożony proces rezerwacji wizyty."),
        ("Observer/Publisher", "BookingEventPublisher i handlery.", "Po rezerwacji reaguje kolejka i SignalR."),
        ("Dependency Injection", "ServiceCollectionExtensions.", "Zależności są podawane przez kontener ASP.NET Core."),
        ("Iterator", "AvailabilitySlotCollection.", "Pozwala przechodzić po slotach godzinowych."),
        ("Strategy/Adapter", "IPhotoStorageService, IEmailSender, IAppointmentBookedHandler.", "Implementacje usług można podmienić bez zmiany kontrolerów."),
    ])
    doc.add_heading("Funkcjonalność", level=1)
    bullets(doc, [
        "Rezerwacja wizyty przechodzi przez fasadę, która waliduje dane i kolizje.",
        "Dane są zapisywane przez repozytorium, a nie bezpośrednio w kontrolerze.",
        "Po utworzeniu wizyty publisher uruchamia handlery kolejki i SignalR.",
        "Fabryka tworzy spójne encje domenowe.",
        "Iterator generuje dostępne godziny pracy salonu.",
        "Dependency Injection scala wzorce i implementacje w jednym miejscu.",
    ])
    add_image(doc, assets["ph_flow"], "Rysunek 2. Miejsce na screen testu przepływu w Swaggerze.", 6.7)
    doc.add_heading("Repository", level=1)
    doc.add_paragraph("Repository oddziela logikę aplikacji od sposobu przechowywania danych. Ten sam kontroler może pracować z Cosmos DB albo repozytorium in-memory, bo zna tylko interfejs.")
    add_image(doc, assets["repo_interface"], "Rysunek 3. Interfejs repozytorium.", 6.7)
    add_image(doc, assets["cosmos_repo"], "Rysunek 4. Implementacja repozytorium dla Cosmos DB.", 6.7)
    doc.add_heading("Factory", level=1)
    doc.add_paragraph("Factory centralizuje tworzenie encji. Dzięki temu kod nie rozrzuca po aplikacji instrukcji new Customer(), new Hairdresser() itd. w miejscach, w których wystarczy wskazać typ encji.")
    add_image(doc, assets["factory"], "Rysunek 5. BookingEntityFactory.", 6.7)
    doc.add_heading("Facade", level=1)
    doc.add_paragraph("Facade upraszcza kontroler Appointments. Zamiast wykonywać wiele kroków w kontrolerze, endpoint przekazuje wizytę do AppointmentBookingFacade, a fasada wykonuje pełen proces.")
    add_image(doc, assets["facade"], "Rysunek 6. Fasada rezerwacji.", 6.7)
    add_image(doc, assets["facade_validation"], "Rysunek 7. Walidacja w fasadzie.", 6.7)
    doc.add_heading("Observer / Publisher", level=1)
    doc.add_paragraph("Po zapisaniu wizyty aplikacja publikuje zdarzenie. Dzięki temu dodanie nowej reakcji, np. SMS albo powiadomienia push, nie wymaga przebudowy fasady.")
    add_image(doc, assets["publisher"], "Rysunek 8. Publisher zdarzenia rezerwacji.", 6.7)
    add_image(doc, assets["queue_handler"], "Rysunek 9. Handler kolejki jako obserwator zdarzenia.", 6.7)
    add_image(doc, assets["signalr_handler"], "Rysunek 10. Handler SignalR jako obserwator zdarzenia.", 6.7)
    doc.add_heading("Dependency Injection", level=1)
    doc.add_paragraph("Dependency Injection pozwala wstrzykiwać interfejsy do kontrolerów i serwisów. To sprawia, że klasy są krótsze, mniej sprzężone i łatwiejsze do testowania.")
    add_image(doc, assets["services_di"], "Rysunek 11. Rejestracja zależności.", 6.7)
    doc.add_heading("Iterator", level=1)
    doc.add_paragraph("AvailabilitySlotCollection implementuje IEnumerable<DateTimeOffset>, więc można używać LINQ do filtrowania wolnych terminów.")
    add_image(doc, assets["iterator"], "Rysunek 12. Iterator dostępnych slotów.", 6.7)
    doc.add_heading("Strategy / Adapter usług", level=1)
    doc.add_paragraph("Interfejsy dla zdjęć, e-maili i handlerów zdarzeń pełnią rolę punktów rozszerzeń. Kod aplikacji nie musi znać szczegółów Azure SDK w kontrolerach.")
    add_image(doc, assets["blob_photo"], "Rysunek 13. Strategia zapisu zdjęć.", 6.7)
    add_image(doc, assets["email_sender"], "Rysunek 14. Strategia wysyłania e-maili.", 6.7)
    doc.add_heading("Przygotowanie projektu", level=1)
    numbered(doc, [
        "Utworzenie modeli domenowych i klasy BookingEntity.",
        "Dodanie interfejsu repozytorium.",
        "Dodanie implementacji Cosmos DB.",
        "Dodanie fabryki encji.",
        "Dodanie fasady rezerwacji.",
        "Dodanie publishera i handlerów zdarzeń.",
        "Rejestracja zależności w kontenerze DI.",
        "Test przepływu w Swaggerze.",
    ])
    add_image(doc, assets["ph_vs"], "Rysunek 15. Miejsce na screen rozwiązania w Visual Studio.", 6.7)
    add_image(doc, assets["repo_interface"], "Rysunek 16. Etap: repozytorium.", 6.7)
    add_image(doc, assets["facade"], "Rysunek 17. Etap: fasada rezerwacji.", 6.7)
    add_image(doc, assets["publisher"], "Rysunek 18. Etap: zdarzenia i obserwatorzy.", 6.7)
    path = OUT / "Raport_Wzorce_HairSalonBooking.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    assets = build_assets()
    print(cloud_report(assets))
    print(patterns_report(assets))
